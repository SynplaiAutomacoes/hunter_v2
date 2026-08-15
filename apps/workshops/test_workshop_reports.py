from __future__ import annotations

from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from djmoney.money import Money
from openpyxl import load_workbook

from apps.accounts.models import Account
from apps.budget.models import Budget, BudgetStatus, BudgetType
from apps.catalog.models.groups import CatalogGroup
from apps.catalog.models.products import Product
from apps.catalog.models.services import Service
from apps.collaborators.models import WorkshopCollaborator, WorkshopMember
from apps.core.infrastructure.excel_report_style import EXCEL_CONTENT_TYPE
from apps.customer.models import Customer, Vehicle
from apps.iam.utils import get_or_create_director_role
from apps.workorder.models import WorkOrder, WorkOrderItem, WorkOrderItemBenefitType, WorkOrderStatus
from apps.workshops.models.workshop_costs import WorkshopCost, WorkshopCostItem
from apps.workshops.models.workshops import Workshop
from apps.workshops.services.workshop_reports import (
    build_approval_rate_report,
    build_mechanic_rework_report,
    build_profitability_report,
    build_warranty_return_report,
)
from apps.workshops.util.monthly_costs import create_default_monthly_costs, get_mechanic_salary_monthly_cost

User = get_user_model()


class WorkshopReportsTests(TestCase):
    def setUp(self) -> None:
        self.account = Account.objects.create(name="Conta Relatorios")
        self.user = User.objects.create_user(username="relatorios-user", password="secret", cpf="52998224725")
        self.user.account = self.account
        self.user.save(update_fields=["account"])
        self.workshop = Workshop.objects.create(
            account=self.account,
            name="Oficina Relatorios",
            cnpj="12.345.678/0001-90",
            phone="+5511999999999",
            address="Rua Relatorio, 1",
        )
        self.other_workshop = Workshop.objects.create(
            account=self.account,
            name="Outra Oficina",
            cnpj="12.345.678/0001-91",
            phone="+5511888888888",
            address="Rua Relatorio, 2",
        )
        role = get_or_create_director_role(account=self.account)
        WorkshopMember.objects.create(user=self.user, workshop=self.workshop, role=role, is_active=True)
        self.client.force_login(self.user)
        session = self.client.session
        session["active_workshop_id"] = self.workshop.pk
        session.save()

        self.customer = Customer.objects.create(
            workshop=self.workshop,
            name="Cliente Relatorio",
            cpf_or_cnpj="11144477735",
            email="cliente@example.invalid",
        )
        self.vehicle = Vehicle.objects.create(
            workshop=self.workshop,
            customer=self.customer,
            plate="ABC1D23",
            brand="Marca",
            model="Modelo",
            year_fabrication="2020",
            year_model="2021",
            color="Prata",
        )
        other_customer = Customer.objects.create(
            workshop=self.other_workshop,
            name="Outro Cliente",
            cpf_or_cnpj="39053344705",
            email="outro@example.invalid",
        )
        self.other_vehicle = Vehicle.objects.create(
            workshop=self.other_workshop,
            customer=other_customer,
            plate="XYZ9K87",
            brand="Marca",
            model="Modelo",
            year_fabrication="2020",
            year_model="2021",
            color="Preto",
        )
        group = CatalogGroup.objects.create(workshop=self.workshop, name="Grupo Relatorio")
        self.product = Product.objects.create(
            workshop=self.workshop,
            group=group,
            code="REL-1",
            name="Peca Relatorio",
            cost_price=Money(80, "BRL"),
            selling_price=Money(120, "BRL"),
        )
        self.mechanic_a = self._create_mechanic(suffix=1, name="Ana Mecanica")
        self.mechanic_b = self._create_mechanic(suffix=2, name="Bruno Mecanico")
        self.july = timezone.make_aware(datetime(2026, 7, 15, 12, 0, 0))

    def _create_mechanic(self, *, suffix: int, name: str) -> WorkshopCollaborator:
        return WorkshopCollaborator.objects.create(
            workshop=self.workshop,
            name=name,
            cpf=f"1234567890{suffix}",
            birth_date=date(1990, 1, 1),
            salary=Money(2000, "BRL"),
            admission_date=date(2025, 1, 1),
            collaborator_type=WorkshopCollaborator.CollaboratorType.PRODUCTIVE,
        )

    def _create_workorder(
        self,
        *,
        workshop: Workshop,
        customer: Customer,
        vehicle: Vehicle,
        budget_type: str,
        delivered_at,
        cost_amount: Decimal = Decimal("50.00"),
        stored_total: Decimal = Decimal("200.00"),
        quantity: int = 1,
        collaborators: list[WorkshopCollaborator] | None = None,
    ) -> WorkOrder:
        budget = Budget.objects.create(
            workshop=workshop,
            customer=customer,
            vehicle=vehicle,
            entry_date=date(2026, 7, 2),
            status=BudgetStatus.DRAFT,
            budget_type=budget_type,
            stored_total_amount=Money(stored_total, "BRL"),
        )
        workorder = WorkOrder.objects.create(
            workshop=workshop,
            budget=budget,
            status=WorkOrderStatus.APPROVED,
            budget_type=budget_type,
            delivered_at=delivered_at,
            stored_total_amount=Money(stored_total, "BRL"),
        )
        WorkOrderItem.objects.create(
            workshop=workshop,
            workorder=workorder,
            product=self.product if workshop.pk == self.workshop.pk else None,
            quantity=quantity,
            product_cost_price=Money(cost_amount, "BRL"),
            product_selling_price=Money(stored_total, "BRL"),
            item_benefit_type=WorkOrderItemBenefitType.WARRANTY if budget_type == "warranty" else WorkOrderItemBenefitType.NORMAL,
        )
        if collaborators:
            workorder.collaborators.set(collaborators)
        return workorder

    def test_rework_ranks_by_loss_and_ignores_other_workshop_and_courtesy(self) -> None:
        self._create_workorder(
            workshop=self.workshop,
            customer=self.customer,
            vehicle=self.vehicle,
            budget_type="warranty",
            delivered_at=self.july,
            quantity=4,
            collaborators=[self.mechanic_a],
        )
        self._create_workorder(
            workshop=self.workshop,
            customer=self.customer,
            vehicle=self.vehicle,
            budget_type="warranty",
            delivered_at=self.july,
            quantity=1,
            collaborators=[self.mechanic_b],
        )
        self._create_workorder(
            workshop=self.workshop,
            customer=self.customer,
            vehicle=self.vehicle,
            budget_type="courtesy",
            delivered_at=self.july,
            cost_amount=Decimal("999.00"),
            collaborators=[self.mechanic_b],
        )
        other_customer = self.other_vehicle.customer
        self._create_workorder(
            workshop=self.other_workshop,
            customer=other_customer,
            vehicle=self.other_vehicle,
            budget_type="warranty",
            delivered_at=self.july,
            cost_amount=Decimal("500.00"),
            collaborators=[self.mechanic_a],
        )

        by_loss = build_mechanic_rework_report(workshop=self.workshop, month=7, year=2026, sort_key="prejuizo")
        self.assertEqual([row.mechanic_name for row in by_loss.rows], ["Ana Mecanica", "Bruno Mecanico"])
        self.assertGreater(by_loss.rows[0].accumulated_loss, by_loss.rows[1].accumulated_loss)
        self.assertEqual(by_loss.rows[0].vehicle_count, 1)

        extra_vehicle = Vehicle.objects.create(
            workshop=self.workshop,
            customer=self.customer,
            plate="DEF4G56",
            brand="Marca",
            model="Outro",
            year_fabrication="2021",
            year_model="2021",
            color="Azul",
        )
        self._create_workorder(
            workshop=self.workshop,
            customer=self.customer,
            vehicle=extra_vehicle,
            budget_type="warranty",
            delivered_at=self.july,
            cost_amount=Decimal("5.00"),
            collaborators=[self.mechanic_b],
        )
        by_vehicles = build_mechanic_rework_report(workshop=self.workshop, month=7, year=2026, sort_key="veiculos")
        self.assertEqual(by_vehicles.rows[0].mechanic_name, "Bruno Mecanico")
        self.assertEqual(by_vehicles.rows[0].vehicle_count, 2)

    def test_profitability_lists_only_delivered_sale_os(self) -> None:
        self._create_workorder(
            workshop=self.workshop,
            customer=self.customer,
            vehicle=self.vehicle,
            budget_type="sale",
            delivered_at=self.july,
            stored_total=Decimal("300.00"),
        )
        self._create_workorder(
            workshop=self.workshop,
            customer=self.customer,
            vehicle=self.vehicle,
            budget_type="warranty",
            delivered_at=self.july,
            stored_total=Decimal("80.00"),
        )
        report = build_profitability_report(workshop=self.workshop, month=7, year=2026)
        self.assertEqual(len(report.rows), 1)
        self.assertEqual(report.rows[0].total_amount, Decimal("120.00"))
        self.assertEqual(report.rows[0].customer_name, "Cliente Relatorio")
        self.assertFalse(callable(report.rows[0].workorder_number))

    def test_profitability_uses_stored_budget_rentability(self) -> None:
        workorder = self._create_workorder(
            workshop=self.workshop,
            customer=self.customer,
            vehicle=self.vehicle,
            budget_type="sale",
            delivered_at=self.july,
        )
        budget = workorder.budget
        assert budget is not None
        Budget.objects.filter(pk=budget.pk).update(stored_rentability=Decimal("66.35"))
        report = build_profitability_report(workshop=self.workshop, month=7, year=2026)
        self.assertEqual(len(report.rows), 1)
        self.assertEqual(report.rows[0].profitability_percent, Decimal("66.35"))

    def test_warranty_return_includes_mechanic_value_and_cost(self) -> None:
        workorder = self._create_workorder(
            workshop=self.workshop,
            customer=self.customer,
            vehicle=self.vehicle,
            budget_type="warranty",
            delivered_at=self.july,
            cost_amount=Decimal("40.00"),
            stored_total=Decimal("90.00"),
            collaborators=[self.mechanic_a],
        )
        service = Service.objects.create(
            workshop=self.workshop,
            name="Servico Relatorio",
            duration=timedelta(0),
            suggested_cost=Money(25, "BRL"),
            selling_price=Money(50, "BRL"),
        )
        WorkOrderItem.objects.create(
            workshop=self.workshop,
            workorder=workorder,
            service=service,
            quantity=1,
            service_cost_price=Money(25, "BRL"),
            service_selling_price=Money(50, "BRL"),
            item_benefit_type=WorkOrderItemBenefitType.WARRANTY,
        )
        report = build_warranty_return_report(workshop=self.workshop, month=7, year=2026)
        self.assertEqual(len(report.rows), 1)
        self.assertEqual(report.rows[0].workorder_id, workorder.pk)
        self.assertEqual(report.rows[0].mechanic_names, "Ana Mecanica")
        self.assertEqual(report.rows[0].total_amount, Decimal("170.00"))
        self.assertEqual(report.rows[0].product_cost_amount, Decimal("80.00"))
        self.assertEqual(report.rows[0].service_cost_amount, Decimal("25.00"))
        self.assertEqual(report.total_product_cost, Decimal("80.00"))
        self.assertEqual(report.total_service_cost, Decimal("25.00"))

    def test_warranty_service_cost_matches_gestor_pdf_mechanic_cost(self) -> None:
        today = timezone.localdate()
        create_default_monthly_costs(workshop=self.workshop)
        workshop_cost = WorkshopCost.objects.create(
            workshop=self.workshop,
            month=today.month,
            year=today.year,
            mechanic_quantity=1,
            work_days_per_month=22,
            working_hours_per_month=Decimal("100.00"),
        )
        mechanic_monthly = get_mechanic_salary_monthly_cost(workshop=self.workshop)
        assert mechanic_monthly is not None
        WorkshopCostItem.objects.create(
            workshop_cost=workshop_cost,
            monthly_cost=mechanic_monthly,
            amount=Money("14900.00", "BRL"),
        )
        workorder = self._create_workorder(
            workshop=self.workshop,
            customer=self.customer,
            vehicle=self.vehicle,
            budget_type="warranty",
            delivered_at=self.july,
            collaborators=[self.mechanic_a],
        )
        service = Service.objects.create(
            workshop=self.workshop,
            name="Servico Mecanico Relatorio",
            duration=timedelta(hours=1),
            suggested_cost=Money(25, "BRL"),
            selling_price=Money(50, "BRL"),
        )
        WorkOrderItem.objects.create(
            workshop=self.workshop,
            workorder=workorder,
            service=service,
            quantity=1,
            duration=timedelta(hours=1),
            service_cost_price=Money(25, "BRL"),
            service_selling_price=Money(50, "BRL"),
            item_benefit_type=WorkOrderItemBenefitType.WARRANTY,
        )
        report = build_warranty_return_report(workshop=self.workshop, month=7, year=2026)
        self.assertEqual(len(report.rows), 1)
        self.assertEqual(report.rows[0].product_cost_amount, Decimal("80.00"))
        self.assertEqual(report.rows[0].service_cost_amount, Decimal("149.00"))

    def test_approval_rate_lists_sale_budgets_with_reasons(self) -> None:
        Budget.objects.create(
            workshop=self.workshop,
            customer=self.customer,
            vehicle=self.vehicle,
            entry_date=date(2026, 7, 3),
            status=BudgetStatus.APPROVED,
            budget_type=BudgetType.SALE,
            stored_total_amount=Money("100.00", "BRL"),
        )
        Budget.objects.create(
            workshop=self.workshop,
            customer=self.customer,
            vehicle=self.vehicle,
            entry_date=date(2026, 7, 4),
            status=BudgetStatus.REJECTED,
            budget_type=BudgetType.SALE,
            rejection_reason="Preco alto",
            stored_total_amount=Money("80.00", "BRL"),
        )
        Budget.objects.create(
            workshop=self.workshop,
            customer=self.customer,
            vehicle=self.vehicle,
            entry_date=date(2026, 7, 5),
            status=BudgetStatus.CANCELLED,
            budget_type=BudgetType.SALE,
            cancellation_reason="Cliente desistiu",
            stored_total_amount=Money("70.00", "BRL"),
        )
        Budget.objects.create(
            workshop=self.workshop,
            customer=self.customer,
            vehicle=self.vehicle,
            entry_date=date(2026, 7, 6),
            status=BudgetStatus.APPROVED,
            budget_type=BudgetType.WARRANTY,
        )
        report = build_approval_rate_report(workshop=self.workshop, month=7, year=2026)
        self.assertEqual(len(report.rows), 3)
        reasons = {row.reason for row in report.rows}
        self.assertIn("Preco alto", reasons)
        self.assertIn("Cliente desistiu", reasons)
        self.assertEqual(report.approved_count, 1)
        self.assertEqual(report.rejected_count, 1)
        self.assertEqual(report.cancelled_count, 1)

        rejected_only = build_approval_rate_report(workshop=self.workshop, month=7, year=2026, status=BudgetStatus.REJECTED)
        self.assertEqual(len(rejected_only.rows), 1)
        self.assertEqual(rejected_only.rows[0].reason, "Preco alto")

    def test_management_pages_and_excel_export(self) -> None:
        self._create_workorder(
            workshop=self.workshop,
            customer=self.customer,
            vehicle=self.vehicle,
            budget_type="warranty",
            delivered_at=self.july,
            collaborators=[self.mechanic_a],
        )
        home = self.client.get(reverse("workshops:workshop_reports_home"), {"mes": 7, "ano": 2026})
        self.assertEqual(home.status_code, 200)
        self.assertContains(home, "Relatórios da Oficina")

        rework = self.client.get(reverse("workshops:workshop_report_rework"), {"mes": 7, "ano": 2026})
        self.assertEqual(rework.status_code, 200)
        self.assertContains(rework, "Ana Mecanica")
        self.assertContains(rework, "Voltar")
        self.assertContains(rework, reverse("workshops:workshop_reports_home"))

        excel = self.client.get(reverse("workshops:workshop_report_rework_excel"), {"mes": 7, "ano": 2026})
        self.assertEqual(excel.status_code, 200)
        self.assertEqual(excel["Content-Type"], EXCEL_CONTENT_TYPE)
        workbook = load_workbook(BytesIO(excel.content))
        values = [value for row in workbook.active.iter_rows(values_only=True) for value in row]
        self.assertTrue(any(isinstance(value, str) and "ANA MECANICA" in value.upper() for value in values if value))

        warranty = self.client.get(reverse("workshops:workshop_report_warranty"), {"mes": 7, "ano": 2026})
        self.assertEqual(warranty.status_code, 200)
        self.assertContains(warranty, "Custo de produtos")
        self.assertContains(warranty, "Custo de serviços")

        warranty_excel = self.client.get(reverse("workshops:workshop_report_warranty_excel"), {"mes": 7, "ano": 2026})
        self.assertEqual(warranty_excel.status_code, 200)
        warranty_workbook = load_workbook(BytesIO(warranty_excel.content))
        warranty_values = [value for row in warranty_workbook.active.iter_rows(values_only=True) for value in row]
        self.assertIn("Custo de produtos", warranty_values)
        self.assertIn("Custo de serviços", warranty_values)
