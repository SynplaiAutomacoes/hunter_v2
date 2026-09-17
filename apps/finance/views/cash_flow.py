from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any
from urllib.parse import urlencode

from django.contrib.auth.mixins import LoginRequiredMixin
from django.db.models import Q
from django.http import HttpResponse
from django.urls import reverse
from django.views.generic import TemplateView, View
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from apps.core.domain.contracts.documents import DocumentRenderRequest
from apps.core.infrastructure.pdf.renderer import build_pdf_http_response, render_template_request_to_pdf

from apps.core.infrastructure.search import apply_text_search
from apps.core.workorder_numbers import format_workorder_reference
from apps.finance.forms.emission_ui import format_money
from apps.finance.models.bank_account import BankAccount
from apps.finance.models.financial_group import FinancialGroup
from apps.finance.models.financial_movement import FinancialMovement
from apps.finance.models.payment_method import PaymentMethod
from apps.finance.services.payroll_visibility import resolve_payroll_movement_display
from apps.finance.services.reports import FinancialOverview, build_financial_overview, filter_grouped_movements_for_reporting
from apps.finance.services.workorder_financial_movements import build_workorder_revenue_description
from apps.workorder.models import WorkOrder, WorkOrderPaymentMethod
from apps.workshops.mixin import WorkshopScopedMixin

_SORTABLE_ATTRS = frozenset({"due_date"})
_DEFAULT_SORT = "-due_date"
_PAGE_SIZE = 40


class CashFlowView(LoginRequiredMixin, WorkshopScopedMixin, TemplateView):
    model = FinancialMovement
    template_name = "finance/cash_flow/cash_flow.html"
    htmx_template_name = "finance/cash_flow/partials/cash_flow_rows.html"
    workshop_permission_codename = "view_financialmovement"

    @staticmethod
    def _resolve_result_tone(value: object) -> str:
        amount = Decimal(str(getattr(value, "amount", value) or 0))
        if amount > 0:
            return "credit"
        if amount < 0:
            return "debit"
        return "neutral"

    @staticmethod
    def _resolve_money_amount(value: object) -> Decimal:
        return Decimal(str(getattr(value, "amount", value) or 0))

    def _parse_date_param(self, raw_value: str | None) -> date | None:
        value = str(raw_value or "").strip()
        if not value:
            return None
        try:
            return date.fromisoformat(value)
        except ValueError:
            return None

    def _resolve_period(self) -> dict[str, Any]:
        start_date = self._parse_date_param(self.request.GET.get("data_inicial"))
        end_date = self._parse_date_param(self.request.GET.get("data_final"))
        return {
            "start_date": start_date,
            "end_date": end_date,
            "period_is_custom": start_date is not None or end_date is not None,
        }

    def _parse_page(self) -> int:
        raw_page = str(self.request.GET.get("page") or "1").strip()
        try:
            page = int(raw_page)
        except ValueError:
            return 1
        return page if page > 0 else 1

    def _parse_sort(self) -> str:
        sort = str(self.request.GET.get("sort") or _DEFAULT_SORT).strip()
        sort_attr = sort.lstrip("-")
        if sort_attr not in _SORTABLE_ATTRS:
            return _DEFAULT_SORT
        if sort not in {sort_attr, f"-{sort_attr}"}:
            return _DEFAULT_SORT
        return sort

    def _get_filter_params(self) -> dict[str, Any]:
        period = self._resolve_period()
        return {
            "start_date": period["start_date"],
            "end_date": period["end_date"],
            "period_is_custom": period["period_is_custom"],
            "agent": self.request.GET.get("agente", "").strip(),
            "payment_method_id": self.request.GET.get("forma_pagamento"),
            "budget_plan_id": self.request.GET.get("plano_orcamentario"),
            "movement_type": self.request.GET.get("tipo_movimentacao"),
            "bank_account_id": self.request.GET.get("conta_bancaria"),
            "search": self.request.GET.get("search", "").strip(),
            "sort": self._parse_sort(),
            "page": self._parse_page(),
        }

    def _build_query_params(self, *, overrides: dict[str, str | None] | None = None, exclude: set[str] | None = None) -> dict[str, str]:
        filter_params = self._get_filter_params()
        params: dict[str, str] = {}

        if filter_params["start_date"]:
            params["data_inicial"] = filter_params["start_date"].isoformat()
        if filter_params["end_date"]:
            params["data_final"] = filter_params["end_date"].isoformat()
        if filter_params["agent"]:
            params["agente"] = filter_params["agent"]
        if filter_params["payment_method_id"]:
            params["forma_pagamento"] = str(filter_params["payment_method_id"])
        if filter_params["budget_plan_id"]:
            params["plano_orcamentario"] = str(filter_params["budget_plan_id"])
        if filter_params["movement_type"]:
            params["tipo_movimentacao"] = str(filter_params["movement_type"])
        if filter_params["bank_account_id"]:
            params["conta_bancaria"] = str(filter_params["bank_account_id"])
        if filter_params["search"]:
            params["search"] = filter_params["search"]
        if filter_params["sort"] and filter_params["sort"] != _DEFAULT_SORT:
            params["sort"] = filter_params["sort"]

        if overrides:
            for key, value in overrides.items():
                if value is None or value == "":
                    params.pop(key, None)
                else:
                    params[key] = value

        if exclude:
            for key in exclude:
                params.pop(key, None)

        return params

    def _build_url(self, *, overrides: dict[str, str | None] | None = None, exclude: set[str] | None = None) -> str:
        params = self._build_query_params(overrides=overrides, exclude=exclude)
        base = reverse("finance:cash_flow")
        if not params:
            return base
        return f"{base}?{urlencode(params)}"

    def _sort_toggle_url(self, *, attr: str) -> str:
        current = self._parse_sort()
        next_sort = attr if current == f"-{attr}" else f"-{attr}"
        return self._build_url(overrides={"sort": next_sort})

    @staticmethod
    def _resolve_workorder_description(workorder: WorkOrder) -> str:
        if getattr(workorder, "budget", None) is None:
            return "-"
        return build_workorder_revenue_description(workorder=workorder)

    def _apply_workorder_payment_aware_date_filter(self, queryset, *, lookup: str, value: date):
        return queryset.filter(
            Q(**{lookup: value})
            | Q(
                movement_kind=FinancialMovement.MovementKind.WORKORDER_PARENT,
                workorder__isnull=False,
                workorder__payments__isnull=False,
                **{f"workorder__payments__{lookup}": value},
            )
        ).distinct()

    def _get_financial_movements_queryset(self):
        queryset = (
            filter_grouped_movements_for_reporting(FinancialMovement.objects.filter(workshop=self.workshop))
            .select_related(
                "source",
                "budget_plan",
                "bank_account",
                "payment_method",
                "workorder",
                "workorder__budget",
                "workorder__budget__customer",
            )
            .prefetch_related("workorder__payments", "workorder__payments__payment_method")
            .order_by("-due_date", "-pk")
        )

        filter_params = self._get_filter_params()

        if filter_params["start_date"]:
            queryset = self._apply_workorder_payment_aware_date_filter(queryset, lookup="due_date__gte", value=filter_params["start_date"])
        if filter_params["end_date"]:
            queryset = self._apply_workorder_payment_aware_date_filter(queryset, lookup="due_date__lte", value=filter_params["end_date"])

        if filter_params["agent"]:
            queryset = apply_text_search(queryset, search_value=filter_params["agent"], lookups=("source__name", "workorder__budget__customer__name"))

        if filter_params["payment_method_id"]:
            pm_filter = Q(payment_method_id=filter_params["payment_method_id"]) | Q(
                movement_kind=FinancialMovement.MovementKind.WORKORDER_PARENT,
                workorder__isnull=False,
                workorder__payments__payment_method_id=filter_params["payment_method_id"],
            )
            queryset = queryset.filter(pm_filter).distinct()

        if filter_params["budget_plan_id"]:
            queryset = queryset.filter(budget_plan_id=filter_params["budget_plan_id"])

        if filter_params["movement_type"]:
            queryset = queryset.filter(direction=filter_params["movement_type"])

        if filter_params["bank_account_id"]:
            if filter_params["bank_account_id"] == "none":
                queryset = queryset.filter(bank_account__isnull=True)
            else:
                queryset = queryset.filter(bank_account_id=filter_params["bank_account_id"])

        if filter_params["search"]:
            queryset = apply_text_search(
                queryset,
                search_value=filter_params["search"],
                lookups=("description", "source__name", "nf_number", "workorder__budget__customer__name"),
            ).distinct()

        return queryset

    def _filter_workorder_payments_for_rows(
        self,
        *,
        payments: list[WorkOrderPaymentMethod],
        filter_start_date: date | None,
        filter_end_date: date | None,
        payment_method_id: str | None,
        bank_account_id: str | None,
    ) -> list[WorkOrderPaymentMethod]:
        reconciled_movement_by_payment_id: dict[int, FinancialMovement] = getattr(self, "_reconciled_workorder_payment_movements", {})
        aggregate_movement_by_workorder_id: dict[int, FinancialMovement] = getattr(self, "_reconciled_aggregate_workorder_movements", {})
        filtered_payments = []
        for payment in payments:
            payment_amount = self._resolve_money_amount(payment.total_paid)
            if payment.due_date is None or payment_amount <= Decimal("0.00"):
                continue
            if filter_start_date and payment.due_date < filter_start_date:
                continue
            if filter_end_date and payment.due_date > filter_end_date:
                continue
            if payment_method_id and str(getattr(payment.payment_method, "pk", "")) != str(payment_method_id):
                continue
            payment_movement = reconciled_movement_by_payment_id.get(payment.pk)
            if payment_movement is None:
                payment_movement = aggregate_movement_by_workorder_id.get(getattr(payment, "workorder_id", None))
            if payment_movement is None:
                continue
            if bank_account_id:
                if bank_account_id == "none":
                    if payment_movement.bank_account_id is not None:
                        continue
                elif str(payment_movement.bank_account_id or "") != str(bank_account_id):
                    continue
            filtered_payments.append(payment)
        return filtered_payments

    def _build_workorder_payment_row(self, *, movement: FinancialMovement, payment: WorkOrderPaymentMethod) -> dict[str, object]:
        workorder = movement.workorder
        reconciled_movement_by_payment_id: dict[int, FinancialMovement] = getattr(self, "_reconciled_workorder_payment_movements", {})
        aggregate_movement_by_workorder_id: dict[int, FinancialMovement] = getattr(self, "_reconciled_aggregate_workorder_movements", {})
        payment_movement = reconciled_movement_by_payment_id.get(payment.pk)
        if payment_movement is None and workorder is not None:
            payment_movement = aggregate_movement_by_workorder_id.get(workorder.pk)
        if payment_movement is None:
            payment_movement = movement
        customer = getattr(getattr(workorder, "budget", None), "customer", None) if workorder is not None else None
        payment_method = getattr(payment_movement, "payment_method", None) or getattr(payment, "payment_method", None)
        agent, description = resolve_payroll_movement_display(movement=payment_movement, user=self.request.user, workshop=self.workshop, request=self.request)

        return {
            "component": f"workorder-payment-{payment.pk}",
            "is_expandable": False,
            "type_badge": payment_movement.report_direction_badge,
            "due_date": payment.due_date,
            "agent": (getattr(customer, "name", "-") or "-") if workorder is not None else agent,
            "origin": format_workorder_reference(workorder) if workorder is not None else "-",
            "description": self._resolve_workorder_description(workorder) if workorder is not None else description,
            "budget_plan": payment_movement.report_budget_plan_display,
            "account": payment_movement.report_bank_account_display,
            "payment_type": getattr(payment_method, "description", "-") or "-",
            "edit_url": reverse("finance:financial_movement_update", kwargs={"pk": payment_movement.pk}),
            "total": {
                "text": f"+ {format_money(payment.total_paid)}",
                "class": "text-success font-semibold whitespace-nowrap",
            },
            "details": [],
        }

    def _build_financial_movement_row(self, movement: FinancialMovement) -> dict[str, object] | None:
        workorder = getattr(movement, "workorder", None)
        customer = getattr(getattr(workorder, "budget", None), "customer", None) if workorder is not None else None

        if movement.movement_kind == FinancialMovement.MovementKind.WORKORDER_PARENT and workorder is not None:
            return None
        if not movement.is_paid or not movement.is_reconciled:
            return None

        agent, description = resolve_payroll_movement_display(movement=movement, user=self.request.user, workshop=self.workshop, request=self.request)

        return {
            "component": f"financial-movement-{movement.pk}",
            "is_expandable": False,
            "type_badge": movement.report_direction_badge,
            "due_date": movement.due_date,
            "agent": agent if not workorder else (getattr(customer, "name", "-") or "-"),
            "origin": movement.report_origin_display if not workorder else format_workorder_reference(workorder),
            "description": description,
            "budget_plan": movement.report_budget_plan_display,
            "account": movement.report_bank_account_display,
            "payment_type": movement.report_payment_method_display,
            "edit_url": reverse("finance:financial_movement_update", kwargs={"pk": movement.pk}),
            "total": movement.report_total_display,
            "details": [],
        }

    def _sort_rows(self, rows: list[dict[str, object]], *, sort: str) -> list[dict[str, object]]:
        sort_attr = sort.lstrip("-")
        reverse = sort.startswith("-")
        if sort_attr == "due_date":
            rows.sort(key=lambda row: (row["due_date"] or date.min, str(row["component"])), reverse=reverse)
        return rows

    def _get_financial_movement_report_rows(self) -> list[dict[str, object]]:
        filter_params = self._get_filter_params()
        rows: list[dict[str, object]] = []
        movements = list(self._get_financial_movements_queryset())
        workorder_ids = sorted({movement.workorder.pk for movement in movements if movement.workorder is not None and movement.movement_kind == FinancialMovement.MovementKind.WORKORDER_PARENT})
        reconciled_workorder_payment_movements = FinancialMovement.objects.none()
        reconciled_aggregate_workorder_movements = FinancialMovement.objects.none()
        if workorder_ids:
            bank_account_id = filter_params["bank_account_id"]
            reconciled_workorder_payment_movements = FinancialMovement.objects.filter(
                workshop=self.workshop,
                movement_kind=FinancialMovement.MovementKind.WORKORDER_PARENT,
                workorder_id__in=workorder_ids,
                workorder_payment__isnull=False,
                is_paid=True,
                is_reconciled=True,
            ).select_related("payment_method", "budget_plan", "bank_account")
            reconciled_aggregate_workorder_movements = FinancialMovement.objects.filter(
                workshop=self.workshop,
                movement_kind=FinancialMovement.MovementKind.WORKORDER_PARENT,
                workorder_id__in=workorder_ids,
                workorder_payment__isnull=True,
                is_paid=True,
                is_reconciled=True,
            ).select_related("payment_method", "budget_plan", "bank_account")
            if bank_account_id:
                if bank_account_id == "none":
                    reconciled_workorder_payment_movements = reconciled_workorder_payment_movements.filter(bank_account__isnull=True)
                    reconciled_aggregate_workorder_movements = reconciled_aggregate_workorder_movements.filter(bank_account__isnull=True)
                else:
                    reconciled_workorder_payment_movements = reconciled_workorder_payment_movements.filter(bank_account_id=bank_account_id)
                    reconciled_aggregate_workorder_movements = reconciled_aggregate_workorder_movements.filter(bank_account_id=bank_account_id)

        self._reconciled_workorder_payment_movements = {movement.workorder_payment.pk: movement for movement in reconciled_workorder_payment_movements if movement.workorder_payment is not None}
        self._reconciled_aggregate_workorder_movements = {
            movement.workorder_id: movement for movement in reconciled_aggregate_workorder_movements if movement.workorder_id is not None
        }

        processed_workorder_ids: set[int] = set()
        for movement in movements:
            workorder = getattr(movement, "workorder", None)
            if workorder is not None and movement.movement_kind == FinancialMovement.MovementKind.WORKORDER_PARENT:
                if workorder.pk in processed_workorder_ids:
                    continue
                processed_workorder_ids.add(workorder.pk)
                payments = list(workorder.payments.all())
                filtered_payments = self._filter_workorder_payments_for_rows(
                    payments=payments,
                    filter_start_date=filter_params["start_date"],
                    filter_end_date=filter_params["end_date"],
                    payment_method_id=filter_params["payment_method_id"],
                    bank_account_id=filter_params["bank_account_id"],
                )
                if filtered_payments:
                    rows.extend(self._build_workorder_payment_row(movement=movement, payment=payment) for payment in filtered_payments)
                continue
            row = self._build_financial_movement_row(movement)
            if row is not None:
                rows.append(row)

        return self._sort_rows(rows, sort=filter_params["sort"])

    def _paginate_rows(self, rows: list[dict[str, object]], *, page: int) -> tuple[list[dict[str, object]], bool]:
        start = (page - 1) * _PAGE_SIZE
        end = start + _PAGE_SIZE
        return rows[start:end], end < len(rows)

    def _build_report_url(self, *, view_name: str, account_id: str | None) -> str:
        params = self._build_query_params(overrides={"conta_bancaria": account_id or None, "page": None})
        base = reverse(view_name)
        if not params:
            return base
        return f"{base}?{urlencode(params)}"

    def _resolve_account_title(self, *, account_id: str | None) -> str:
        if not account_id:
            return "Todas as contas"
        if account_id == "none":
            return "Sem Vínculo"
        try:
            account = BankAccount.objects.get(pk=account_id, workshop=self.workshop)
        except (BankAccount.DoesNotExist, ValueError, TypeError):
            return "Conta bancária"
        return str(account)

    def _build_full_report_context(self) -> dict[str, Any]:
        filter_params = self._get_filter_params()
        account_id = str(filter_params.get("bank_account_id") or "")
        rows = self._get_financial_movement_report_rows()
        overview = build_financial_overview(**self._build_overview_kwargs(filter_params, bank_account_id=account_id or None))
        account_title = self._resolve_account_title(account_id=account_id or None)
        start_date = filter_params["start_date"]
        end_date = filter_params["end_date"]
        if start_date or end_date:
            start_label = start_date.strftime("%d/%m/%Y") if start_date else "…"
            end_label = end_date.strftime("%d/%m/%Y") if end_date else "…"
            period_label = f"{start_label} a {end_label}"
        else:
            period_label = "Todo o período"

        return {
            "workshop": self.workshop,
            "account_title": account_title,
            "report_title": f"Fluxo de contas — {account_title}",
            "period_label": period_label,
            "filter_start_date": start_date,
            "filter_end_date": end_date,
            "financial_movement_report_rows": rows,
            "record_count": len(rows),
            "total_value": format_money(overview.confirmed_result),
            "total_tone": self._resolve_result_tone(overview.confirmed_result),
            "pdf_url": self._build_report_url(view_name="finance:cash_flow_report_pdf", account_id=account_id or None),
            "excel_url": self._build_report_url(view_name="finance:cash_flow_report_excel", account_id=account_id or None),
        }

    def _build_account_card(self, *, name: str, account_id: str, overview: FinancialOverview, selected_account_id: str | None) -> dict[str, object]:
        return {
            "name": name,
            "account_id": account_id,
            "value": format_money(overview.confirmed_result),
            "tone": self._resolve_result_tone(overview.confirmed_result),
            "url": self._build_url(
                overrides={
                    "conta_bancaria": account_id or None,
                    "data_inicial": None,
                    "data_final": None,
                    "page": None,
                }
            ),
            "is_selected": (selected_account_id or "") == account_id,
        }

    def _build_overview_kwargs(self, filter_params: dict[str, Any], *, bank_account_id: str | None) -> dict[str, Any]:
        return {
            "workshop": self.workshop,
            "start_date": filter_params["start_date"],
            "end_date": filter_params["end_date"],
            "search": filter_params["search"],
            "direction": filter_params["movement_type"],
            "paid_status": "paid",
            "reconciliation_status": "reconciled",
            "budget_plan_ids": [filter_params["budget_plan_id"]] if filter_params["budget_plan_id"] else None,
            "bank_account_id": bank_account_id or None,
            "agent": filter_params["agent"],
            "payment_method_id": filter_params["payment_method_id"],
        }

    def _has_active_filters(self, filter_params: dict[str, Any]) -> bool:
        return bool(filter_params["agent"] or filter_params["payment_method_id"] or filter_params["budget_plan_id"] or filter_params["movement_type"] or filter_params["search"] or filter_params["period_is_custom"] or filter_params["bank_account_id"] or filter_params["sort"] != _DEFAULT_SORT)

    def get_template_names(self) -> list[str]:
        if bool(getattr(self.request, "htmx", False)) and self._parse_page() > 1:
            return [self.htmx_template_name]
        return [self.template_name]

    def get_context_data(self, **kwargs: Any) -> dict[str, Any]:
        context = super().get_context_data(**kwargs)
        filter_params = self._get_filter_params()
        bank_accounts = list(BankAccount.objects.filter(workshop=self.workshop).order_by("bank_name"))
        selected_account_id = str(filter_params.get("bank_account_id") or "")

        account_cards = [
            self._build_account_card(
                name="Todas as contas",
                account_id="",
                overview=build_financial_overview(**self._build_overview_kwargs(filter_params, bank_account_id=None)),
                selected_account_id=selected_account_id,
            )
        ]
        for account in bank_accounts:
            account_cards.append(
                self._build_account_card(
                    name=str(account),
                    account_id=str(account.pk),
                    overview=build_financial_overview(**self._build_overview_kwargs(filter_params, bank_account_id=str(account.pk))),
                    selected_account_id=selected_account_id,
                )
            )

        all_rows = self._get_financial_movement_report_rows()
        page = filter_params["page"]
        page_rows, has_next_page = self._paginate_rows(all_rows, page=page)
        sort = filter_params["sort"]
        sort_attr = sort.lstrip("-")

        context["account_cards"] = account_cards
        context["filter_start_date"] = filter_params["start_date"]
        context["filter_end_date"] = filter_params["end_date"]
        context["period_is_custom"] = filter_params["period_is_custom"]
        context["financial_movement_report_rows"] = page_rows
        context["has_next_page"] = has_next_page
        context["next_page_url"] = self._build_url(overrides={"page": str(page + 1)}) if has_next_page else None
        context["clear_filters_url"] = reverse("finance:cash_flow")
        context["has_active_filters"] = self._has_active_filters(filter_params)
        context["preserved_query_params"] = self._build_query_params(exclude={"page"})
        context["sort"] = sort
        context["sort_due_date_url"] = self._sort_toggle_url(attr="due_date")
        context["sort_due_date_is_asc"] = sort_attr == "due_date" and not sort.startswith("-")
        context["sort_due_date_is_desc"] = sort_attr == "due_date" and sort.startswith("-")
        context["bank_accounts"] = bank_accounts
        context["payment_methods"] = PaymentMethod.objects.filter(workshop=self.workshop).order_by("description")
        context["budget_plans"] = FinancialGroup.objects.filter(workshop=self.workshop).order_by("sort_key")
        context["movement_types"] = FinancialMovement.MovementDirection.choices
        return context


class CashFlowReportModalView(CashFlowView):
    template_name = "finance/cash_flow/partials/cash_flow_report_modal.html"

    def get_template_names(self) -> list[str]:
        return [self.template_name]

    def get_context_data(self, **kwargs: Any) -> dict[str, Any]:
        context = super(TemplateView, self).get_context_data(**kwargs)
        context.update(self._build_full_report_context())
        context["is_pdf"] = False
        return context


class CashFlowReportPdfView(LoginRequiredMixin, WorkshopScopedMixin, View):
    model = FinancialMovement
    workshop_permission_codename = "view_financialmovement"

    def get(self, request, *args: Any, **kwargs: Any) -> HttpResponse:
        report_view = CashFlowView()
        report_view.setup(request)
        report_view.request = request
        report_view.workshop = self.workshop
        context = report_view._build_full_report_context()
        context["is_pdf"] = True
        document = render_template_request_to_pdf(
            DocumentRenderRequest(
                template_name="finance/cash_flow/pdf/cash_flow_report.html",
                context=context,
                filename=f"fluxo_de_contas_{context['account_title']}.pdf",
            )
        )
        return build_pdf_http_response(document=document, download=True)


class CashFlowReportExcelView(LoginRequiredMixin, WorkshopScopedMixin, View):
    model = FinancialMovement
    workshop_permission_codename = "view_financialmovement"

    def get(self, request, *args: Any, **kwargs: Any) -> HttpResponse:
        report_view = CashFlowView()
        report_view.setup(request)
        report_view.request = request
        report_view.workshop = self.workshop
        context = report_view._build_full_report_context()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Fluxo de contas"
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(fill_type="solid", fgColor="1E3A8A")
        headers = ["Tipo", "Data", "Agente", "Origem", "Descrição", "Plano Orçamentário", "Conta Bancária", "Tipo Pagamento", "Valor"]
        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=column, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        sheet.cell(row=2, column=1, value="Conta")
        sheet.cell(row=2, column=2, value=str(context["account_title"]))
        sheet.cell(row=3, column=1, value="Período")
        sheet.cell(row=3, column=2, value=str(context["period_label"]))
        sheet.cell(row=4, column=1, value="Total")
        sheet.cell(row=4, column=2, value=str(context["total_value"]))

        start_row = 6
        for index, row in enumerate(context["financial_movement_report_rows"], start=start_row):
            due_date = row.get("due_date")
            total = row.get("total") or {}
            sheet.cell(row=index, column=1, value=str((row.get("type_badge") or {}).get("text") or "-"))
            sheet.cell(row=index, column=2, value=due_date.strftime("%d/%m/%Y") if hasattr(due_date, "strftime") else "-")
            sheet.cell(row=index, column=3, value=str(row.get("agent") or "-"))
            sheet.cell(row=index, column=4, value=str(row.get("origin") or "-"))
            sheet.cell(row=index, column=5, value=str(row.get("description") or "-"))
            sheet.cell(row=index, column=6, value=str(row.get("budget_plan") or "-"))
            sheet.cell(row=index, column=7, value=str(row.get("account") or "-"))
            sheet.cell(row=index, column=8, value=str(row.get("payment_type") or "-"))
            sheet.cell(row=index, column=9, value=str(total.get("text") or "-"))

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        filename = f"fluxo_de_contas_{context['account_title']}.xlsx".replace(" ", "_")
        response = HttpResponse(buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        response["X-Content-Type-Options"] = "nosniff"
        response["Cache-Control"] = "no-store"
        return response
