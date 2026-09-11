from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch
from uuid import uuid4

from django.test import RequestFactory, SimpleTestCase, TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.collaborators.test_commissions import create_workshop
from apps.finance.models.finance import (
    NfeItem,
    NfeRequest,
    NfeRequestStatus,
    NfseItem,
    NfseRequest,
    NfseRequestStatus,
    StandaloneNfeLine,
    StandaloneNfseLine,
)
from apps.finance.services.issued_documents_report import (
    build_issued_documents_report,
    build_issued_documents_report_excel,
)
from apps.finance.views.issued_documents_report import IssuedDocumentsReportExcelView, IssuedDocumentsReportPdfView


class IssuedDocumentsReportServiceTests(TestCase):
    def setUp(self) -> None:
        self.workshop = create_workshop(suffix=91)
        self.other_workshop = create_workshop(suffix=92)

    def _create_standalone_nfe(
        self,
        *,
        workshop,
        number: str,
        status: str,
        amount: Decimal,
        recipient_name: str = "Cliente NF-e",
        created_at: datetime | None = None,
    ) -> NfeRequest:
        nfe_request = NfeRequest.objects.create(
            workshop=workshop,
            workorder=None,
            status=status,
            reserved_number=int(number),
            recipient_name=recipient_name,
        )
        StandaloneNfeLine.objects.create(
            nfe_request=nfe_request,
            description="Peca avulsa",
            product_code="P1",
            ncm="12345678",
            quantity=Decimal("1"),
            unit_value=amount,
        )
        NfeItem.objects.create(
            workshop=workshop,
            request=nfe_request,
            uuid=uuid4(),
            number=number,
            status="aprovado" if status == NfeRequestStatus.APPROVED else "cancelado",
        )
        if created_at is not None:
            NfeRequest.objects.filter(pk=nfe_request.pk).update(criado_em=created_at)
            nfe_request.refresh_from_db()
        return nfe_request

    def _create_standalone_nfse(
        self,
        *,
        workshop,
        number: str,
        status: str,
        amount: Decimal,
        recipient_name: str = "Cliente NFS-e",
        created_at: datetime | None = None,
    ) -> NfseRequest:
        nfse_request = NfseRequest.objects.create(
            workshop=workshop,
            workorder=None,
            status=status,
            reserved_rps_number=int(number),
            recipient_name=recipient_name,
        )
        StandaloneNfseLine.objects.create(
            nfse_request=nfse_request,
            description="Servico avulso",
            quantity=Decimal("1"),
            unit_value=amount,
        )
        NfseItem.objects.create(
            workshop=workshop,
            request=nfse_request,
            uuid=uuid4(),
            number=number,
            rps_number=number,
            status="aprovado" if status == NfseRequestStatus.APPROVED else "cancelado",
        )
        if created_at is not None:
            NfseRequest.objects.filter(pk=nfse_request.pk).update(criado_em=created_at)
            nfse_request.refresh_from_db()
        return nfse_request

    def test_approved_nfe_enters_total_and_canceled_does_not(self) -> None:
        self._create_standalone_nfe(
            workshop=self.workshop,
            number="1001",
            status=NfeRequestStatus.APPROVED,
            amount=Decimal("150.00"),
        )
        self._create_standalone_nfe(
            workshop=self.workshop,
            number="1002",
            status=NfeRequestStatus.CANCELED,
            amount=Decimal("80.00"),
            recipient_name="Cliente Cancelado",
        )
        self._create_standalone_nfe(
            workshop=self.workshop,
            number="1003",
            status=NfeRequestStatus.REPROVED,
            amount=Decimal("99.00"),
            recipient_name="Cliente Reprovado",
        )

        report = build_issued_documents_report(workshop=self.workshop, note_type="nfe")

        self.assertEqual(report.record_count, 2)
        self.assertEqual({row.number for row in report.rows}, {"1001", "1002"})
        self.assertEqual(report.total_amount, Decimal("150.00"))
        canceled = next(row for row in report.rows if row.number == "1002")
        self.assertFalse(canceled.include_in_total)
        self.assertEqual(canceled.amount, Decimal("80.00"))

    def test_nfse_report_does_not_mix_nfe(self) -> None:
        self._create_standalone_nfe(
            workshop=self.workshop,
            number="2001",
            status=NfeRequestStatus.APPROVED,
            amount=Decimal("200.00"),
        )
        self._create_standalone_nfse(
            workshop=self.workshop,
            number="3001",
            status=NfseRequestStatus.APPROVED,
            amount=Decimal("45.50"),
        )

        report = build_issued_documents_report(workshop=self.workshop, note_type="nfse")

        self.assertEqual(report.record_count, 1)
        self.assertEqual(report.rows[0].number, "3001")
        self.assertEqual(report.total_amount, Decimal("45.50"))

    def test_period_filter_excludes_notes_outside_range(self) -> None:
        inside = timezone.make_aware(datetime(2026, 8, 15, 10, 0))
        outside = timezone.make_aware(datetime(2026, 7, 15, 10, 0))
        self._create_standalone_nfe(
            workshop=self.workshop,
            number="4001",
            status=NfeRequestStatus.APPROVED,
            amount=Decimal("10.00"),
            created_at=inside,
        )
        self._create_standalone_nfe(
            workshop=self.workshop,
            number="4002",
            status=NfeRequestStatus.APPROVED,
            amount=Decimal("20.00"),
            created_at=outside,
        )

        report = build_issued_documents_report(
            workshop=self.workshop,
            note_type="nfe",
            start_date=date(2026, 8, 1),
            end_date=date(2026, 8, 31),
        )

        self.assertEqual(report.record_count, 1)
        self.assertEqual(report.rows[0].number, "4001")
        self.assertEqual(report.total_amount, Decimal("10.00"))
        self.assertIn("01/08/2026", report.period_label)

    def test_other_workshop_notes_are_excluded(self) -> None:
        self._create_standalone_nfe(
            workshop=self.workshop,
            number="5001",
            status=NfeRequestStatus.APPROVED,
            amount=Decimal("30.00"),
        )
        self._create_standalone_nfe(
            workshop=self.other_workshop,
            number="5002",
            status=NfeRequestStatus.APPROVED,
            amount=Decimal("999.00"),
        )

        report = build_issued_documents_report(workshop=self.workshop, note_type="nfe")

        self.assertEqual(report.record_count, 1)
        self.assertEqual(report.rows[0].number, "5001")
        self.assertEqual(report.total_amount, Decimal("30.00"))

    def test_excel_contains_total_amount(self) -> None:
        self._create_standalone_nfe(
            workshop=self.workshop,
            number="6001",
            status=NfeRequestStatus.APPROVED,
            amount=Decimal("123.45"),
            recipient_name="Cliente Excel",
        )
        report = build_issued_documents_report(workshop=self.workshop, note_type="nfe")
        document = build_issued_documents_report_excel(report=report)

        self.assertTrue(document.filename.endswith(".xlsx"))
        workbook = load_workbook(BytesIO(document.content))
        sheet = workbook.active
        values = [cell.value for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row) for cell in row]
        self.assertIn("Cliente Excel", values)
        self.assertIn(123.45, values)
        self.assertTrue(any(isinstance(value, str) and value.startswith("Valor total") for value in values))


class IssuedDocumentsReportViewTests(TestCase):
    def setUp(self) -> None:
        self.factory = RequestFactory()
        self.workshop = create_workshop(suffix=93)
        self.user = SimpleNamespace(is_authenticated=True, is_active=True)

    def _create_approved_nfe(self) -> NfeRequest:
        nfe_request = NfeRequest.objects.create(
            workshop=self.workshop,
            workorder=None,
            status=NfeRequestStatus.APPROVED,
            reserved_number=7001,
            recipient_name="Cliente View",
        )
        StandaloneNfeLine.objects.create(
            nfe_request=nfe_request,
            description="Peca",
            product_code="P1",
            ncm="12345678",
            quantity=Decimal("1"),
            unit_value=Decimal("50.00"),
        )
        return nfe_request

    def test_excel_view_returns_download(self) -> None:
        self._create_approved_nfe()
        request = self.factory.get(reverse("finance:issued_documents_report_excel"), {"tipo": "nfe"})
        request.user = self.user
        view = IssuedDocumentsReportExcelView()
        view.setup(request)
        view.request = request
        view.workshop = self.workshop

        response = view.get(request)

        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertGreater(len(response.content), 0)

    @patch("apps.finance.services.issued_documents_report.render_template_request_to_pdf")
    def test_pdf_view_returns_download(self, render_pdf_mock) -> None:
        self._create_approved_nfe()
        render_pdf_mock.return_value = SimpleNamespace(
            content=b"%PDF-1.4 mock",
            filename="relatorio_nfe.pdf",
            content_type="application/pdf",
        )
        request = self.factory.get(reverse("finance:issued_documents_report_pdf"), {"tipo": "nfe"})
        request.user = self.user
        view = IssuedDocumentsReportPdfView()
        view.setup(request)
        view.request = request
        view.workshop = self.workshop

        response = view.get(request)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertEqual(response.content, b"%PDF-1.4 mock")
        render_pdf_mock.assert_called_once()

    def test_missing_note_type_returns_400(self) -> None:
        request = self.factory.get(reverse("finance:issued_documents_report_excel"))
        request.user = self.user
        view = IssuedDocumentsReportExcelView()
        view.setup(request)
        view.request = request
        view.workshop = self.workshop

        response = view.get(request)

        self.assertEqual(response.status_code, 400)
        self.assertIn(b"NF-e ou NFS-e", response.content)


class IssuedDocumentsReportTemplateTests(SimpleTestCase):
    def test_central_shows_report_button_outside_fiscal_operation(self) -> None:
        from django.template import Context
        from django.template.loader import get_template

        request = RequestFactory().get("/finance/notas-emitidas/")
        request.user = SimpleNamespace(is_authenticated=True)
        template = get_template("finance/issued_documents_list.html").template
        html = str(
            template.render(
                Context(
                    {
                        "request": request,
                        "filter_state": {
                            "has_selected_period": False,
                            "selected_note_type": "all",
                            "search_raw": "",
                            "start_raw": "",
                            "end_raw": "",
                        },
                        "note_type_choices": (("all", "Todas"), ("nfe", "Nota Fiscal de Produto"), ("nfse", "Nota Fiscal Serviço")),
                        "issued_note_rows": [],
                        "issued_notes_total": 0,
                        "issued_nfe_total": 0,
                        "issued_nfse_total": 0,
                        "download_xml_url": "/finance/notas-emitidas/download/xml/",
                        "download_pdfs_url": "/finance/notas-emitidas/download/pdfs/",
                        "report_pdf_url": "/finance/notas-emitidas/relatorio/pdf/",
                        "report_excel_url": "/finance/notas-emitidas/relatorio/excel/",
                        "fiscal_operation": "",
                        "fiscal_operation_label": "",
                        "fiscal_operation_continuation_label": "",
                        "selected_fiscal_nfe": None,
                        "fiscal_selection_reset_url": "",
                    }
                )
            )
        )

        self.assertIn("Relatório", html)
        self.assertIn("issued-documents-report-modal", html)
        self.assertIn('data-report-base-url="/finance/notas-emitidas/relatorio/pdf/"', html)
        self.assertIn('data-report-base-url="/finance/notas-emitidas/relatorio/excel/"', html)
        self.assertIn("report-export-spinner", html)
        self.assertIn('data-report-export="true"', html)

