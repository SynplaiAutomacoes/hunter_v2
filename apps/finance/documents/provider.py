from __future__ import annotations

from io import BytesIO
import re
import unicodedata
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.core.documents.contract import DocumentPayload, DocumentRenderRequest
from apps.core.documents.renderer import render_template_request_to_pdf
from apps.workshops.models.workshops import Workshop


_DRE_EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_PRIMARY_FILL = PatternFill(fill_type="solid", fgColor="1E3A8A")
_PRIMARY_SOFT_FILL = PatternFill(fill_type="solid", fgColor="DBEAFE")
_SECONDARY_FILL = PatternFill(fill_type="solid", fgColor="EFF6FF")
_RESULT_FILL = PatternFill(fill_type="solid", fgColor="E0E7FF")
_POSITIVE_FILL = PatternFill(fill_type="solid", fgColor="DCFCE7")
_NEGATIVE_FILL = PatternFill(fill_type="solid", fgColor="FEE2E2")
_BORDER_COLOR = "BFDBFE"
_THIN_BORDER = Border(
    left=Side(style="thin", color=_BORDER_COLOR),
    right=Side(style="thin", color=_BORDER_COLOR),
    top=Side(style="thin", color=_BORDER_COLOR),
    bottom=Side(style="thin", color=_BORDER_COLOR),
)
_CURRENCY_FORMAT = "R$ #,##0.00;[Red]-R$ #,##0.00"
_DATE_FORMAT = "DD/MM/YYYY"


def build_dre_pdf_render_request(*, context: dict[str, object], request=None, filename: str | None = None) -> DocumentRenderRequest:
    selected_workshop = context.get("selected_workshop")
    workshop = selected_workshop if isinstance(selected_workshop, Workshop) else None
    data_inicial_label = str(context.get("data_inicial_label") or "-")
    data_final_label = str(context.get("data_final_label") or "-")
    resolved_filename = filename or _build_filename(
        workshop=workshop,
        data_inicial_label=data_inicial_label,
        data_final_label=data_final_label,
        extension="pdf",
    )

    render_context = dict(context)
    render_context["request"] = request

    return DocumentRenderRequest(
        template_name="finance/dre/pdf/visualizarPDF.html",
        context=render_context,
        filename=resolved_filename,
    )


def render_dre_pdf_document(*, context: dict[str, object], request=None, filename: str | None = None) -> DocumentPayload:
    render_request = build_dre_pdf_render_request(context=context, request=request, filename=filename)
    return render_template_request_to_pdf(render_request)


def build_dre_excel_document(*, context: dict[str, object], filename: str | None = None) -> DocumentPayload:
    workbook = Workbook()
    summary_sheet = workbook.active
    if summary_sheet is None:
        raise ValueError("Nao foi possivel inicializar a planilha de resumo do DRE.")
    summary_sheet.title = "Resumo"
    details_sheet = workbook.create_sheet("Detalhes")

    _populate_summary_sheet(summary_sheet=summary_sheet, context=context)
    _populate_details_sheet(details_sheet=details_sheet, context=context)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    selected_workshop = context.get("selected_workshop")
    workshop = selected_workshop if isinstance(selected_workshop, Workshop) else None
    data_inicial_label = str(context.get("data_inicial_label") or "-")
    data_final_label = str(context.get("data_final_label") or "-")
    resolved_filename = filename or _build_filename(workshop=workshop, data_inicial_label=data_inicial_label, data_final_label=data_final_label, extension="xlsx")
    return DocumentPayload(content=buffer.getvalue(), filename=resolved_filename, content_type=_DRE_EXCEL_CONTENT_TYPE)


def _build_filename(*, workshop: Workshop | None, data_inicial_label: str, data_final_label: str, extension: str) -> str:
    workshop_fragment = _normalize_filename_fragment(workshop.name) if workshop is not None else "consolidado"
    start_fragment = _normalize_filename_fragment(data_inicial_label)
    end_fragment = _normalize_filename_fragment(data_final_label)
    return f"dre_{workshop_fragment}_{start_fragment}_{end_fragment}.{extension}"


def _populate_summary_sheet(*, summary_sheet, context: dict[str, object]) -> None:
    selected_workshop = context.get("selected_workshop")
    workshop = selected_workshop if isinstance(selected_workshop, Workshop) else None
    summary_cards = _get_context_dict_list(context, "dre_summary_cards")
    dre_rows = _get_context_dict_list(context, "dre_rows")
    summary_sheet.sheet_view.showGridLines = False
    summary_sheet.freeze_panes = "A7"
    summary_sheet.merge_cells("A1:C1")
    title_cell = summary_sheet["A1"]
    title_cell.value = "DRE - Demonstracao do Resultado do Exercicio"
    title_cell.fill = _PRIMARY_FILL
    title_cell.font = Font(color="FFFFFF", bold=True, size=16)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    summary_sheet.row_dimensions[1].height = 28

    metadata_rows = [
        ("Filial", workshop.name if workshop is not None else "Consolidado"),
        ("Periodo", f"{context.get('data_inicial_label', '-')} ate {context.get('data_final_label', '-')}"),
    ]
    for index, (label, value) in enumerate(metadata_rows, start=3):
        summary_sheet[f"A{index}"] = label
        summary_sheet[f"B{index}"] = value
        summary_sheet[f"A{index}"].font = Font(bold=True, color="1E3A8A")
        summary_sheet[f"A{index}"].fill = _SECONDARY_FILL
        summary_sheet[f"A{index}"].border = _THIN_BORDER
        summary_sheet[f"B{index}"].border = _THIN_BORDER
        summary_sheet[f"B{index}"].alignment = Alignment(wrap_text=True)

    summary_sheet["A7"] = "Resumo"
    summary_sheet["B7"] = "Valor"
    for cell_ref in ("A7", "B7"):
        summary_sheet[cell_ref].fill = _PRIMARY_SOFT_FILL
        summary_sheet[cell_ref].font = Font(bold=True, color="1E3A8A")
        summary_sheet[cell_ref].border = _THIN_BORDER
        summary_sheet[cell_ref].alignment = Alignment(horizontal="center")

    for index, card in enumerate(summary_cards, start=8):
        summary_sheet[f"A{index}"] = str(card.get("label") or "-")
        summary_sheet[f"B{index}"] = _as_excel_number(card.get("amount"))
        summary_sheet[f"A{index}"].fill = _SECONDARY_FILL
        summary_sheet[f"B{index}"].fill = _SECONDARY_FILL
        summary_sheet[f"A{index}"].font = Font(bold=True)
        summary_sheet[f"B{index}"].font = Font(bold=True, color="1E3A8A")
        summary_sheet[f"B{index}"].number_format = _CURRENCY_FORMAT
        summary_sheet[f"B{index}"].alignment = Alignment(horizontal="right")
        summary_sheet[f"A{index}"].border = _THIN_BORDER
        summary_sheet[f"B{index}"].border = _THIN_BORDER

    table_start = max(8, 8 + len(summary_cards)) + 2
    headers = ["Descricao", "Formula", "Valor"]
    for column_index, header in enumerate(headers, start=1):
        cell = summary_sheet.cell(row=table_start, column=column_index, value=header)
        cell.fill = _PRIMARY_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = _THIN_BORDER

    for row_index, row in enumerate(dre_rows, start=table_start + 1):
        amount = _as_excel_number(row.get("amount"))
        tone = str(row.get("tone") or "")
        fill = _resolve_row_fill(tone=tone)
        values = [
            str(row.get("label") or "-"),
            str(row.get("formula") or "-"),
            amount,
        ]
        for column_index, value in enumerate(values, start=1):
            cell = summary_sheet.cell(row=row_index, column=column_index, value=value)
            cell.border = _THIN_BORDER
            cell.fill = fill
            if column_index == 3:
                cell.number_format = _CURRENCY_FORMAT
                cell.alignment = Alignment(horizontal="right")
                cell.font = Font(bold=tone in {"highlight", "result"})
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.font = Font(bold=tone in {"highlight", "result"})

    _set_column_widths(summary_sheet, {"A": 42, "B": 48, "C": 18})


def _populate_details_sheet(*, details_sheet, context: dict[str, object]) -> None:
    dre_rows = _get_context_dict_list(context, "dre_rows")
    details_sheet.sheet_view.showGridLines = False
    details_sheet.freeze_panes = "A4"
    details_sheet.merge_cells("A1:F1")
    title_cell = details_sheet["A1"]
    title_cell.value = "Detalhamento do DRE"
    title_cell.fill = _PRIMARY_FILL
    title_cell.font = Font(color="FFFFFF", bold=True, size=15)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    details_sheet.row_dimensions[1].height = 26

    details_sheet["A2"] = "Periodo"
    details_sheet["B2"] = f"{context.get('data_inicial_label', '-')} ate {context.get('data_final_label', '-')}"
    details_sheet["A2"].font = Font(bold=True, color="1E3A8A")

    headers = ["Linha DRE", "Resumo", "Referencia", "Data Entrada", "Data Saida", "Valor"]
    for column_index, header in enumerate(headers, start=1):
        cell = details_sheet.cell(row=3, column=column_index, value=header)
        cell.fill = _PRIMARY_SOFT_FILL
        cell.font = Font(bold=True, color="1E3A8A")
        cell.alignment = Alignment(horizontal="center")
        cell.border = _THIN_BORDER

    next_row = 4
    for dre_row in dre_rows:
        details_value = dre_row.get("details")
        details = details_value if isinstance(details_value, list) else []
        if not details:
            continue
        tone = str(dre_row.get("tone") or "")
        fill = _resolve_row_fill(tone=tone)
        for detail in details:
            details_sheet.cell(row=next_row, column=1, value=str(dre_row.get("label") or "-")).fill = fill
            details_sheet.cell(row=next_row, column=2, value=str(detail.get("summary") or "-")).fill = fill
            details_sheet.cell(row=next_row, column=3, value=str(detail.get("reference") or "-")).fill = fill
            entry_date_cell = details_sheet.cell(row=next_row, column=4, value=detail.get("entry_date"))
            payment_date_cell = details_sheet.cell(row=next_row, column=5, value=detail.get("payment_date"))
            amount_cell = details_sheet.cell(row=next_row, column=6, value=_as_excel_number(detail.get("amount")))
            entry_date_cell.fill = fill
            payment_date_cell.fill = fill
            amount_cell.fill = fill
            amount_cell.number_format = _CURRENCY_FORMAT
            amount_cell.alignment = Alignment(horizontal="right")
            if entry_date_cell.value:
                entry_date_cell.number_format = _DATE_FORMAT
            if payment_date_cell.value:
                payment_date_cell.number_format = _DATE_FORMAT
            for column_index in range(1, 7):
                details_sheet.cell(row=next_row, column=column_index).border = _THIN_BORDER
                if column_index != 6:
                    details_sheet.cell(row=next_row, column=column_index).alignment = Alignment(vertical="center", wrap_text=True)
            next_row += 1

    if next_row == 4:
        details_sheet.merge_cells("A4:F4")
        empty_cell = details_sheet["A4"]
        empty_cell.value = "Nao ha detalhes disponiveis para os filtros informados."
        empty_cell.fill = _SECONDARY_FILL
        empty_cell.alignment = Alignment(horizontal="center")
        empty_cell.font = Font(italic=True, color="1E3A8A")
        empty_cell.border = _THIN_BORDER

    _set_column_widths(details_sheet, {"A": 34, "B": 44, "C": 18, "D": 14, "E": 14, "F": 18})


def _resolve_row_fill(*, tone: str) -> PatternFill:
    if tone == "positive":
        return _POSITIVE_FILL
    if tone == "negative":
        return _NEGATIVE_FILL
    if tone in {"highlight", "result"}:
        return _RESULT_FILL
    return _SECONDARY_FILL


def _as_excel_number(value: object) -> float:
    amount: Any = getattr(value, "amount") if hasattr(value, "amount") else value
    if isinstance(amount, Decimal):
        return float(amount)
    if isinstance(amount, (int, float)):
        return float(amount)
    if isinstance(amount, str):
        return float(amount or 0)
    return 0.0


def _get_context_dict_list(context: dict[str, object], key: str) -> list[dict[str, object]]:
    value = context.get(key)
    if not isinstance(value, list):
        return []
    return [item for item in value if isinstance(item, dict)]


def _set_column_widths(worksheet, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        worksheet.column_dimensions[get_column_letter(worksheet[f"{column}1"].column)].width = width


def _normalize_filename_fragment(value: str) -> str:
    normalized_value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    normalized_value = normalized_value.lower().strip()
    normalized_value = re.sub(r"[^a-z0-9]+", "_", normalized_value)
    normalized_value = normalized_value.strip("_")
    return normalized_value or "-"
