from __future__ import annotations

import io
from datetime import date, timedelta
from decimal import Decimal

from django.test import SimpleTestCase, TestCase
from djmoney.money import Money
from openpyxl import Workbook

from apps.catalog.importing import KitImportAdapter, ProductImportAdapter, ServiceImportAdapter
from apps.catalog.models.groups import CatalogGroup
from apps.catalog.models.kits import Kit, KitProduct, KitService
from apps.catalog.models.products import Product
from apps.catalog.models.services import Service
from apps.core.infrastructure.importing.kit_report import expand_kit_report, looks_like_kit_report
from apps.core.infrastructure.importing.mapper import apply_mapping, coerce_value, parse_decimal, parse_duration, suggest_mapping
from apps.core.infrastructure.importing.schema import FieldType, ImportEntitySchema, ImportField, ImportRowError, ParsedSpreadsheet
from apps.core.infrastructure.importing.session_store import create_import_session, get_import_session
from apps.core.infrastructure.importing.spreadsheet import SpreadsheetParseError, parse_csv_bytes, parse_xls_bytes, parse_xlsx_bytes
from apps.customer.importing import CustomerImportAdapter
from apps.customer.models import Customer
from apps.finance.importing import PayableImportAdapter
from apps.finance.models import FinancialGroup, PaymentMethod
from apps.finance.models.financial_movement import FinancialMovement
from apps.stock.models import StockProduct
from apps.suppliers.importing import SupplierImportAdapter
from apps.suppliers.models import Supplier
from apps.workshops.models.workshops import Workshop


def create_workshop(*, suffix: int = 1) -> Workshop:
    return Workshop.objects.create(
        name=f"Oficina Import {suffix}",
        cnpj=f"11.222.333/0001-{suffix:02d}",
        phone="+5511999999999",
        address="Rua Teste, 123",
    )


class SpreadsheetParseTests(SimpleTestCase):
    def test_parse_csv_with_semicolon_delimiter(self) -> None:
        content = "Nome;Email\nAna;ana@example.com\n".encode("utf-8")
        parsed = parse_csv_bytes(content, filename="clientes.csv")
        self.assertEqual(parsed.headers, ["Nome", "Email"])
        self.assertEqual(len(parsed.rows), 1)
        self.assertEqual(parsed.rows[0]["Nome"], "Ana")

    def test_parse_xlsx_headers_and_rows(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.append(["Código", "Nome"])
        sheet.append(["P1", "Peça A"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        parsed = parse_xlsx_bytes(buffer.getvalue(), filename="produtos.xlsx")
        self.assertEqual(parsed.headers, ["Código", "Nome"])
        self.assertEqual(parsed.rows[0]["Código"], "P1")

    def test_parse_xls_headers_and_rows(self) -> None:
        import xlwt  # type: ignore[import-untyped]

        workbook = xlwt.Workbook()
        sheet = workbook.add_sheet("Dados")
        sheet.write(0, 0, "Código")
        sheet.write(0, 1, "Nome")
        sheet.write(1, 0, "P1")
        sheet.write(1, 1, "Peça A")
        buffer = io.BytesIO()
        workbook.save(buffer)
        parsed = parse_xls_bytes(buffer.getvalue(), filename="produtos.xls")
        self.assertEqual(parsed.source_format, "xls")
        self.assertEqual(parsed.headers, ["Código", "Nome"])
        self.assertEqual(parsed.rows[0]["Código"], "P1")
        self.assertEqual(parsed.rows[0]["Nome"], "Peça A")

    def test_reject_empty_csv(self) -> None:
        with self.assertRaises(SpreadsheetParseError):
            parse_csv_bytes(b"", filename="vazio.csv")


class MapperTests(SimpleTestCase):
    def test_suggest_mapping_matches_labels(self) -> None:
        schema = ImportEntitySchema(
            entity_key="demo",
            label="Demo",
            fields=(
                ImportField("name", "Nome", required=True),
                ImportField("email", "E-mail", required=True, field_type=FieldType.EMAIL),
            ),
        )
        mapping = suggest_mapping(["nome", "E-mail"], schema)
        self.assertEqual(mapping["name"], "nome")
        self.assertEqual(mapping["email"], "E-mail")

    def test_parse_brazilian_decimal_and_duration(self) -> None:
        self.assertEqual(parse_decimal("1.234,56"), Decimal("1234.56"))
        self.assertEqual(parse_duration("1:30"), timedelta(hours=1, minutes=30))
        self.assertEqual(parse_duration("90"), timedelta(minutes=90))

    def test_parse_duration_accepts_hh_mm_ss_and_ignores_seconds(self) -> None:
        self.assertEqual(parse_duration("00:00:00"), timedelta(0))
        self.assertEqual(parse_duration("01:00:00"), timedelta(hours=1))
        self.assertEqual(parse_duration("1:30:45"), timedelta(hours=1, minutes=30))
        self.assertEqual(parse_duration("02:15:59.999"), timedelta(hours=2, minutes=15))

    def test_apply_mapping_allows_ignored_required_fields(self) -> None:
        schema = ImportEntitySchema(
            entity_key="demo",
            label="Demo",
            fields=(
                ImportField("name", "Nome", required=True),
                ImportField("cost_price", "Custo", required=True, field_type=FieldType.MONEY),
            ),
        )
        mapped = apply_mapping({"Nome": "Peça"}, mapping={"name": "Nome"}, schema=schema)
        self.assertEqual(mapped, {"name": "Peça"})
        self.assertNotIn("cost_price", mapped)

    def test_money_invalid_or_empty_becomes_zero(self) -> None:
        money_field = ImportField("cost_price", "Custo", required=True, field_type=FieldType.MONEY)
        self.assertEqual(coerce_value("", money_field), Decimal("0"))
        self.assertEqual(coerce_value("abc", money_field), Decimal("0"))
        self.assertEqual(coerce_value("10,50", money_field), Decimal("10.50"))


class ImportSessionStoreTests(TestCase):
    def test_create_and_get_session_persists_in_database(self) -> None:
        from django.contrib.auth import get_user_model

        workshop = create_workshop(suffix=90)
        user = get_user_model().objects.create_user(username="import-session", password="x")
        spreadsheet = ParsedSpreadsheet(
            headers=["Nome"],
            rows=[{"Nome": "A"}],
            filename="a.csv",
            source_format="csv",
        )
        token = create_import_session(
            entity_key="customers",
            workshop_id=int(workshop.pk),
            user_id=int(user.pk),
            spreadsheet=spreadsheet,
        )
        payload = get_import_session(token)
        self.assertIsNotNone(payload)
        assert payload is not None
        self.assertEqual(payload["filename"], "a.csv")
        self.assertEqual(payload["rows"], [{"Nome": "A"}])


class CustomerImportAdapterTests(TestCase):
    def test_creates_customer_and_rejects_duplicate(self) -> None:
        workshop = create_workshop(suffix=1)
        adapter = CustomerImportAdapter()
        mapped = {
            "customer_type": "PF",
            "name": "Cliente Importado",
            "cpf_or_cnpj": "52998224725",
            "email": "cliente@import.test",
        }
        created_id = adapter.create_row(workshop, mapped)
        self.assertIsNotNone(created_id)
        self.assertTrue(Customer.objects.filter(pk=created_id, workshop=workshop).exists())

        with self.assertRaises(ImportRowError):
            adapter.create_row(workshop, mapped)


class SupplierImportAdapterTests(TestCase):
    def test_creates_supplier(self) -> None:
        workshop = create_workshop(suffix=2)
        adapter = SupplierImportAdapter()
        created_id = adapter.create_row(
            workshop,
            {"name": "Fornecedor Import", "cnpj": "11444777000161", "email": "f@example.com"},
        )
        self.assertTrue(Supplier.objects.filter(pk=created_id, workshop=workshop).exists())


class ProductImportAdapterTests(TestCase):
    def test_creates_product_with_stock_adjustment(self) -> None:
        workshop = create_workshop(suffix=3)
        CatalogGroup.objects.create(workshop=workshop, name="Geral")
        adapter = ProductImportAdapter()
        mapped = {
            "code": "IMP-01",
            "name": "Produto Importado",
            "unit": "UND",
            "group": "Geral",
            "cost_price": Decimal("10.00"),
            "selling_price": Decimal("20.00"),
            "stock_quantity": Decimal("15"),
            "stock_minimum": 2,
        }
        product_id = adapter.create_row(workshop, mapped, options={"include_stock": True})
        product = Product.objects.get(pk=product_id)
        stock = StockProduct.objects.get(product=product)
        self.assertEqual(stock.current_quantity, Decimal("15"))
        self.assertEqual(stock.minimum_quantity, 2)

    def test_creates_product_with_default_group_and_zero_prices(self) -> None:
        workshop = create_workshop(suffix=31)
        adapter = ProductImportAdapter()
        product_id = adapter.create_row(
            workshop,
            {
                "code": "IMP-02",
                "name": "Produto Sem Grupo",
            },
        )
        product = Product.objects.get(pk=product_id)
        self.assertEqual(product.group.name, "Importação")
        self.assertEqual(product.unit, Product.Unit.UND)
        self.assertEqual(product.cost_price.amount, Decimal("0.00"))
        self.assertEqual(product.selling_price.amount, Decimal("0.00"))

    def test_allows_duplicate_product_names_with_different_codes(self) -> None:
        workshop = create_workshop(suffix=32)
        CatalogGroup.objects.create(workshop=workshop, name="Geral")
        adapter = ProductImportAdapter()
        first = adapter.create_row(
            workshop,
            {"code": "F1", "name": "FILTRO DE AR", "unit": "UND", "group": "Geral", "cost_price": Decimal("1"), "selling_price": Decimal("2")},
        )
        second = adapter.create_row(
            workshop,
            {"code": "F2", "name": "FILTRO DE AR", "unit": "UND", "group": "Geral", "cost_price": Decimal("1"), "selling_price": Decimal("2")},
        )
        self.assertNotEqual(first, second)
        self.assertEqual(Product.objects.filter(workshop=workshop, name__iexact="FILTRO DE AR").count(), 2)

    def test_accepts_long_name_description_and_code(self) -> None:
        """Nome/descrição são TextField; código/sku até 255 (não mais varchar(50))."""
        workshop = create_workshop(suffix=33)
        CatalogGroup.objects.create(workshop=workshop, name="Geral")
        adapter = ProductImportAdapter()
        long_name = "ÓLEO DO MOTOR SW30 DIESEL " + ("X" * 80)
        long_description = "Febi 06161 - Fluido Hidráulico " + ("Y" * 120)
        long_code = "COD-" + ("Z" * 60)
        product_id = adapter.create_row(
            workshop,
            {
                "code": long_code,
                "name": long_name,
                "description": long_description,
                "unit": "UND",
                "group": "Geral",
                "cost_price": Decimal("1"),
                "selling_price": Decimal("2"),
            },
        )
        product = Product.objects.get(pk=product_id)
        self.assertEqual(product.code, long_code)
        self.assertGreater(len(product.code), 50)
        self.assertGreater(len(product.name), 50)
        self.assertGreater(len(product.description), 50)
        self.assertTrue(product.name.casefold().startswith("óleo do motor sw30 diesel"))
        self.assertTrue(product.description.casefold().startswith("febi 06161 - fluido hidráulico"))
        self.assertIn("X" * 80, product.name)
        self.assertIn("Y" * 120, product.description)


class ServiceImportAdapterTests(TestCase):
    def test_creates_service(self) -> None:
        workshop = create_workshop(suffix=4)
        adapter = ServiceImportAdapter()
        created_id = adapter.create_row(
            workshop,
            {
                "name": "Alinhamento Import",
                "duration": timedelta(hours=1),
                "selling_price": Decimal("80.00"),
            },
        )
        self.assertTrue(Service.objects.filter(pk=created_id, workshop=workshop).exists())


class KitReportExpandTests(SimpleTestCase):
    def test_expands_hierarchical_ultracar_report(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.append(["Código Interno", "Nome", "Descrição", "Quantidade de Itens", None, None, None])
        sheet.append(["25", "KIT FREIO", "Kit completo", "2", None, None, None])
        sheet.append([None, "Código Interno", "Descrição", "Custo", "Venda", "Quantidade", "Tipo"])
        sheet.append([None, "P-001", "PASTILHA", "0,00", "0,00", "0.00", "Peça"])
        sheet.append([None, "168", "TROCA PASTILHA", "0,00", "0,00", "1.00", "Peça"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        parsed = parse_xlsx_bytes(buffer.getvalue(), filename="kits.xlsx")
        self.assertTrue(looks_like_kit_report(parsed))
        expanded = expand_kit_report(parsed)
        self.assertEqual(len(expanded.rows), 2)
        self.assertEqual(expanded.rows[0]["Nome do kit"], "KIT FREIO")
        self.assertEqual(expanded.rows[0]["Código ou nome do item"], "P-001")
        self.assertEqual(expanded.rows[0]["Quantidade"], "1")  # 0 → 1
        self.assertEqual(expanded.rows[1]["Código ou nome do item"], "168")
        self.assertEqual(expanded.rows[1]["Quantidade"], "1")


class KitImportAdapterTests(TestCase):
    def test_creates_kit_from_hierarchical_report_rows(self) -> None:
        workshop = create_workshop(suffix=51)
        group = CatalogGroup.objects.create(workshop=workshop, name="Geral")
        Product.objects.create(
            workshop=workshop,
            code="P-001",
            name="PASTILHA",
            unit=Product.Unit.UND,
            group=group,
            cost_price=Money("5.00", "BRL"),
            selling_price=Money("10.00", "BRL"),
        )
        Service.objects.create(
            workshop=workshop,
            name="TROCA PASTILHA",
            duration=timedelta(minutes=30),
            selling_price=Money("50.00", "BRL"),
        )
        adapter = KitImportAdapter()
        batch = adapter.import_rows(
            workshop,
            [
                (
                    2,
                    {
                        "name": "KIT FREIO",
                        "description": "Kit completo",
                        "item_type": "product",
                        "item_ref": "P-001",
                        "item_name": "PASTILHA",
                        "quantity": 0,
                    },
                ),
                (
                    3,
                    {
                        "name": "KIT FREIO",
                        "description": "Kit completo",
                        "item_type": "product",
                        "item_ref": "168",
                        "item_name": "TROCA PASTILHA",
                        "quantity": 1,
                    },
                ),
            ],
        )
        self.assertEqual(batch.success_count, 2, msg="; ".join(r.message for r in batch.results if not r.success))
        kit = Kit.objects.get(workshop=workshop, name__iexact="KIT FREIO")
        self.assertEqual(KitProduct.objects.filter(kit=kit).count(), 1)
        self.assertEqual(KitService.objects.filter(kit=kit).count(), 1)

    def test_skips_missing_items_and_creates_kit_with_found_ones(self) -> None:
        workshop = create_workshop(suffix=52)
        group = CatalogGroup.objects.create(workshop=workshop, name="Geral")
        product = Product.objects.create(
            workshop=workshop,
            code="OK-1",
            name="Produto Ok",
            unit=Product.Unit.UND,
            group=group,
            cost_price=Money("5.00", "BRL"),
            selling_price=Money("10.00", "BRL"),
        )
        adapter = KitImportAdapter()
        batch = adapter.import_rows(
            workshop,
            [
                (2, {"name": "Kit Parcial", "item_type": "product", "item_ref": "OK-1", "quantity": 1}),
                (3, {"name": "Kit Parcial", "item_type": "product", "item_ref": "MISSING", "item_name": "Não Existe", "quantity": 1}),
                (4, {"name": "Kit Parcial", "item_type": "service", "item_ref": "Serviço Fantasma", "quantity": 1}),
            ],
        )
        self.assertEqual(batch.success_count, 1)
        self.assertEqual(batch.error_count, 2)
        kit = Kit.objects.get(workshop=workshop, name__iexact="Kit Parcial")
        self.assertTrue(KitProduct.objects.filter(kit=kit, product=product).exists())
        self.assertEqual(KitService.objects.filter(kit=kit).count(), 0)
        skipped_messages = [r.message for r in batch.results if not r.success]
        self.assertTrue(any("ignorado" in msg.casefold() for msg in skipped_messages))

    def test_creates_kit_from_multiple_rows(self) -> None:
        workshop = create_workshop(suffix=5)
        group = CatalogGroup.objects.create(workshop=workshop, name="Geral")
        product = Product.objects.create(
            workshop=workshop,
            code="KP1",
            name="Peça Kit",
            unit=Product.Unit.UND,
            group=group,
            cost_price=Money("5.00", "BRL"),
            selling_price=Money("10.00", "BRL"),
        )
        service = Service.objects.create(
            workshop=workshop,
            name="Serviço Kit",
            duration=timedelta(minutes=30),
            selling_price=Money("50.00", "BRL"),
        )
        adapter = KitImportAdapter()
        batch = adapter.import_rows(
            workshop,
            [
                (2, {"name": "Kit Freio", "item_type": "product", "item_ref": "KP1", "quantity": 2}),
                (3, {"name": "Kit Freio", "item_type": "service", "item_ref": "Serviço Kit", "quantity": 1}),
            ],
        )
        self.assertEqual(
            batch.success_count,
            2,
            msg="; ".join(f"{r.row_number}:{r.message}" for r in batch.results if not r.success),
        )
        kit = Kit.objects.filter(workshop=workshop).order_by("-pk").first()
        self.assertIsNotNone(kit)
        assert kit is not None
        self.assertEqual(kit.name.casefold(), "kit freio")
        self.assertTrue(KitProduct.objects.filter(kit=kit, product=product, quantity=2).exists())
        self.assertTrue(kit.kit_services.filter(service=service).exists())


class PayableImportAdapterTests(TestCase):
    def test_creates_debit_movement(self) -> None:
        workshop = create_workshop(suffix=6)
        PaymentMethod.objects.create(workshop=workshop, description="Pix")
        FinancialGroup.objects.create(workshop=workshop, name="Despesas Gerais")
        Supplier.objects.create(workshop=workshop, name="Fornecedor Conta")

        adapter = PayableImportAdapter()
        created_id = adapter.create_row(
            workshop,
            {
                "description": "Conta importada",
                "gross_amount": Decimal("150.00"),
                "due_date": date(2026, 10, 20),
                "payment_method": "Pix",
                "budget_plan": "Despesas Gerais",
                "supplier": "Fornecedor Conta",
            },
        )
        movement = FinancialMovement.objects.get(pk=created_id)
        self.assertEqual(movement.direction, FinancialMovement.MovementDirection.DEBIT)
        self.assertEqual(movement.amount.amount, Decimal("150.00"))
