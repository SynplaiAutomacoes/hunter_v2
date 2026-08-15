from __future__ import annotations

from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from djmoney.money import Money
from openpyxl import load_workbook

from apps.accounts.models import Account
from apps.budget.models import Budget, BudgetStatus, BudgetType
from apps.collaborators.models import WorkshopMember
from apps.core.infrastructure.excel_report_style import EXCEL_CONTENT_TYPE
from apps.core.infrastructure.services.dashboard_report_excel import build_dashboard_financial_report_excel
from apps.customer.models import Customer, Vehicle
from apps.iam.utils import get_or_create_director_role
from apps.workshops.models.workshops import Workshop

User = get_user_model()


class DashboardReportExcelTests(TestCase):
    def setUp(self) -> None:
        self.account = Account.objects.create(name="Conta Dashboard Excel")
        self.user = User.objects.create_user(username="dashboard-excel-user", password="secret", cpf="52998224725")
        self.user.account = self.account
        self.user.save(update_fields=["account"])
        self.workshop = Workshop.objects.create(
            account=self.account,
            name="Oficina Dashboard Excel",
            cnpj="12.345.678/0001-90",
            phone="+5511999999999",
            address="Rua Dashboard, 1",
        )
        role = get_or_create_director_role(account=self.account)
        WorkshopMember.objects.create(user=self.user, workshop=self.workshop, role=role, is_active=True)
        self.client.force_login(self.user)
        session = self.client.session
        session["active_workshop_id"] = self.workshop.pk
        session.save()
        self.customer = Customer.objects.create(
            workshop=self.workshop,
            name="Cliente Dashboard",
            cpf_or_cnpj="11144477735",
            email="cliente@example.invalid",
        )
        self.vehicle = Vehicle.objects.create(
            workshop=self.workshop,
            customer=self.customer,
            plate="ABC1D23",
            brand="Marca",
            model="Modelo",
            year_fabrication="2020",
            year_model="2021",
            color="Prata",
        )

    def test_dashboard_excel_export_for_existing_indicator(self) -> None:
        Budget.objects.create(
            workshop=self.workshop,
            customer=self.customer,
            vehicle=self.vehicle,
            entry_date=date(2026, 7, 8),
            status=BudgetStatus.REJECTED,
            budget_type=BudgetType.SALE,
            rejection_reason="Sem peca",
            stored_total_amount=Money("55.00", "BRL"),
        )
        response = self.client.get(
            reverse("core:dashboard_financial_report_excel"),
            {"indicador": "reprovados", "mes": 7, "ano": 2026},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], EXCEL_CONTENT_TYPE)
        workbook = load_workbook(BytesIO(response.content))
        values = [value for row in workbook.active.iter_rows(values_only=True) for value in row]
        self.assertTrue(any(isinstance(value, str) and "REPROVADOS" in value.upper() for value in values if value))
        self.assertIn("Sem peca", values)

    def test_dashboard_excel_builder_writes_hunter_header(self) -> None:
        document = build_dashboard_financial_report_excel(
            context={
                "workshop": self.workshop,
                "indicator": "reprovados",
                "report_title": "Total Reprovados",
                "periodo_label": "Julho de 2026",
                "record_count": 0,
                "is_budget_report": True,
                "value_column_label": "Valor",
                "report_rows": [],
            }
        )
        self.assertTrue(document.filename.endswith(".xlsx"))
        self.assertEqual(document.content_type, EXCEL_CONTENT_TYPE)
        workbook = load_workbook(BytesIO(document.content))
        self.assertTrue(str(workbook.active["A1"].value).startswith("RELATÓRIO"))
