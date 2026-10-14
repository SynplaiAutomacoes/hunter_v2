from __future__ import annotations

import csv
import io
from datetime import date, datetime
from pathlib import Path
from typing import Any, BinaryIO

import xlrd  # type: ignore[import-untyped]
from openpyxl import load_workbook  # type: ignore[import-untyped]
from xlrd import xldate  # type: ignore[import-untyped]

from apps.core.infrastructure.importing.schema import ParsedSpreadsheet

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5_000
ALLOWED_EXTENSIONS = {".csv", ".xlsx", ".xls"}


class SpreadsheetParseError(ValueError):
    """Raised when an uploaded spreadsheet cannot be parsed."""


def detect_extension(filename: str) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise SpreadsheetParseError("Formato inválido. Envie apenas arquivos .xlsx, .xls ou .csv.")
    return suffix


def _normalize_header(value: object, index: int) -> str:
    text = str(value or "").strip()
    if not text:
        return f"coluna_{index + 1}"
    return text


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time() == datetime.min.time():
            return value.date().strftime("%d/%m/%Y")
        return value.strftime("%d/%m/%Y %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_csv_bytes(content: bytes, *, filename: str) -> ParsedSpreadsheet:
    for encoding in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise SpreadsheetParseError("Não foi possível ler o CSV (encoding inválido).")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
        if sample.count(";") > sample.count(","):
            dialect.delimiter = ";"

    reader = csv.reader(io.StringIO(text), dialect)
    try:
        raw_headers = next(reader)
    except StopIteration as exc:
        raise SpreadsheetParseError("O arquivo CSV está vazio.") from exc

    headers = [_normalize_header(h, i) for i, h in enumerate(raw_headers)]
    rows: list[dict[str, Any]] = []
    for row_index, values in enumerate(reader, start=1):
        if row_index > MAX_ROWS:
            raise SpreadsheetParseError(f"O arquivo excede o limite de {MAX_ROWS} linhas.")
        if not any(str(v or "").strip() for v in values):
            continue
        padded = list(values) + [""] * max(0, len(headers) - len(values))
        rows.append({headers[i]: _cell_to_str(padded[i]) for i in range(len(headers))})

    if not rows:
        raise SpreadsheetParseError("O arquivo não contém linhas de dados.")

    return ParsedSpreadsheet(headers=headers, rows=rows, filename=filename, source_format="csv")


def parse_xlsx_bytes(content: bytes, *, filename: str) -> ParsedSpreadsheet:
    workbook = load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            raise SpreadsheetParseError("A planilha está vazia.")

        rows_iter = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows_iter)
        except StopIteration as exc:
            raise SpreadsheetParseError("A planilha está vazia.") from exc

        headers = [_normalize_header(h, i) for i, h in enumerate(raw_headers or ())]

        rows: list[dict[str, Any]] = []
        for row_index, values in enumerate(rows_iter, start=1):
            if row_index > MAX_ROWS:
                raise SpreadsheetParseError(f"O arquivo excede o limite de {MAX_ROWS} linhas.")
            values_list = list(values or ())
            if not any(v is not None and str(v).strip() for v in values_list):
                continue
            padded = values_list + [None] * max(0, len(headers) - len(values_list))
            rows.append({headers[i]: _cell_to_str(padded[i]) for i in range(len(headers))})

        if not rows:
            raise SpreadsheetParseError("A planilha não contém linhas de dados.")

        return ParsedSpreadsheet(headers=headers, rows=rows, filename=filename, source_format="xlsx")
    finally:
        workbook.close()


def _xls_cell_to_str(book: Any, sheet: Any, row_index: int, col_index: int) -> str:
    cell_type = sheet.cell_type(row_index, col_index)
    cell_value = sheet.cell_value(row_index, col_index)

    if cell_type == xlrd.XL_CELL_EMPTY:
        return ""
    if cell_type == xlrd.XL_CELL_DATE:
        try:
            date_tuple = xldate.xldate_as_tuple(cell_value, book.datemode)
            if date_tuple[3:] == (0, 0, 0):
                return date(*date_tuple[:3]).strftime("%d/%m/%Y")
            return datetime(*date_tuple).strftime("%d/%m/%Y %H:%M:%S")
        except (OverflowError, ValueError, xlrd.XLDateError):
            return _cell_to_str(cell_value)
    if cell_type == xlrd.XL_CELL_BOOLEAN:
        return "sim" if cell_value else "não"
    if cell_type == xlrd.XL_CELL_NUMBER:
        return _cell_to_str(float(cell_value))
    return _cell_to_str(cell_value)


def parse_xls_bytes(content: bytes, *, filename: str) -> ParsedSpreadsheet:
    try:
        book = xlrd.open_workbook(file_contents=content)
    except xlrd.XLRDError as exc:
        raise SpreadsheetParseError("Não foi possível ler o arquivo .xls.") from exc

    if book.nsheets < 1:
        raise SpreadsheetParseError("A planilha está vazia.")

    sheet = book.sheet_by_index(0)
    if sheet.nrows < 1:
        raise SpreadsheetParseError("A planilha está vazia.")

    headers = [_normalize_header(_xls_cell_to_str(book, sheet, 0, col), col) for col in range(sheet.ncols)]
    rows: list[dict[str, Any]] = []
    for row_index in range(1, sheet.nrows):
        data_row_number = row_index  # 1-based data index after header
        if data_row_number > MAX_ROWS:
            raise SpreadsheetParseError(f"O arquivo excede o limite de {MAX_ROWS} linhas.")
        values = [_xls_cell_to_str(book, sheet, row_index, col) for col in range(sheet.ncols)]
        if not any(value.strip() for value in values):
            continue
        rows.append({headers[col]: values[col] if col < len(values) else "" for col in range(len(headers))})

    if not rows:
        raise SpreadsheetParseError("A planilha não contém linhas de dados.")

    return ParsedSpreadsheet(headers=headers, rows=rows, filename=filename, source_format="xls")


def parse_uploaded_file(uploaded_file: BinaryIO, *, filename: str) -> ParsedSpreadsheet:
    suffix = detect_extension(filename)
    content = uploaded_file.read()
    if len(content) > MAX_FILE_BYTES:
        raise SpreadsheetParseError(f"O arquivo excede o limite de {MAX_FILE_BYTES // (1024 * 1024)} MB.")
    if not content:
        raise SpreadsheetParseError("O arquivo está vazio.")

    if suffix == ".csv":
        return parse_csv_bytes(content, filename=filename)
    if suffix == ".xls":
        return parse_xls_bytes(content, filename=filename)
    return parse_xlsx_bytes(content, filename=filename)
