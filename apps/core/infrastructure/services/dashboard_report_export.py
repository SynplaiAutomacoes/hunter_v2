from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.core.domain.contracts.documents import DocumentPayload
from apps.core.infrastructure.services.dashboard_query_service import resolve_decimal_amount


EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_HEADER_FILL = PatternFill(fill_type="solid", fgColor="17365D")
_SUBHEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")
_TOTAL_FILL = PatternFill(fill_type="solid", fgColor="D9EAD3")
_BORDER_SIDE = Side(style="thin", color="B7C9D6")
_BORDER = Border(left=_BORDER_SIDE, right=_BORDER_SIDE, top=_BORDER_SIDE, bottom=_BORDER_SIDE)
_CURRENCY_FORMAT = 'R$ #,##0.00'


def build_dashboard_financial_report_excel(*, context: dict[str, Any]) -> DocumentPayload:
    """Build a spreadsheet from the same rows and amounts rendered in the report."""
    workbook = Workbook()
    worksheet = workbook.active
    if worksheet is None:
        raise ValueError("Não foi possível inicializar a planilha do relatório.")

    worksheet.title = "Total Vendido"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A7"
    worksheet.merge_cells("A1:E1")
    worksheet["A1"] = str(context["report_title"]).upper()
    worksheet["A1"].fill = _HEADER_FILL
    worksheet["A1"].font = Font(color="FFFFFF", bold=True, size=15)
    worksheet["A1"].alignment = Alignment(horizontal="center")
    worksheet.row_dimensions[1].height = 28

    metadata = (
        ("Oficina", getattr(context["workshop"], "pdf_name", None) or getattr(context["workshop"], "name", "-")),
        ("Período", context["periodo_label"]),
        ("Total vendido", resolve_decimal_amount(context["total_value"])),
        ("Quantidade de O.S.", context["record_count"]),
    )
    for row_index, (label, value) in enumerate(metadata, start=3):
        worksheet.cell(row=row_index, column=1, value=label).font = Font(bold=True)
        worksheet.cell(row=row_index, column=1).fill = _SUBHEADER_FILL
        worksheet.cell(row=row_index, column=1).border = _BORDER
        value_cell = worksheet.cell(row=row_index, column=2, value=float(value) if label == "Total vendido" else value)
        value_cell.border = _BORDER
        if label == "Total vendido":
            value_cell.number_format = _CURRENCY_FORMAT

    headers = ("O.S.", "Cliente", "Veículo", "Data", "Valor vendido")
    for column_index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=6, column=column_index, value=header)
        cell.fill = _HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = _BORDER

    row_index = 7
    for group in context["workorder_groups"]:
        for workorder in (group.primary_item, *group.child_items):
            amount = resolve_decimal_amount(getattr(workorder, "dashboard_report_amount", 0))
            values = (
                workorder.get_id,
                str(workorder.budget.customer) if workorder.budget and workorder.budget.customer else "-",
                str(workorder.budget.vehicle) if workorder.budget and workorder.budget.vehicle else "-",
                workorder.delivered_at.date() if workorder.delivered_at else workorder.criado_em.date(),
                float(amount),
            )
            for column_index, value in enumerate(values, start=1):
                cell = worksheet.cell(row=row_index, column=column_index, value=value)
                cell.border = _BORDER
                if row_index % 2:
                    cell.fill = PatternFill(fill_type="solid", fgColor="F5F9FC")
            worksheet.cell(row=row_index, column=4).number_format = "DD/MM/YYYY"
            worksheet.cell(row=row_index, column=5).number_format = _CURRENCY_FORMAT
            row_index += 1

    worksheet.cell(row=row_index, column=4, value="TOTAL GERAL").font = Font(bold=True)
    total_cell = worksheet.cell(row=row_index, column=5, value=float(resolve_decimal_amount(context["total_value"])))
    total_cell.font = Font(bold=True)
    total_cell.number_format = _CURRENCY_FORMAT
    for column_index in range(1, 6):
        worksheet.cell(row=row_index, column=column_index).fill = _TOTAL_FILL
        worksheet.cell(row=row_index, column=column_index).border = _BORDER

    for column_index, width in enumerate((14, 30, 36, 14, 20), start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    filename = "relatorio_total_vendido.xlsx"
    return DocumentPayload(content=buffer.getvalue(), filename=filename, content_type=EXCEL_CONTENT_TYPE)
