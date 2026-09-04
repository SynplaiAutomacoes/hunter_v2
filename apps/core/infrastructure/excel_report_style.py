from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Literal

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

from apps.core.domain.contracts.documents import DocumentPayload


EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CURRENCY_FORMAT = r"R$ #,##0.00"
DATE_FORMAT = "DD/MM/YYYY"
PERCENT_FORMAT = "0.0%"

FILL_TITLE = PatternFill(fill_type="solid", fgColor="0D2137")
FILL_SUBTITLE = PatternFill(fill_type="solid", fgColor="2E5FA3")
FILL_HEADER = PatternFill(fill_type="solid", fgColor="1A3C6E")
FILL_ZEBRA = PatternFill(fill_type="solid", fgColor="F0F5FF")
FILL_WHITE = PatternFill(fill_type="solid", fgColor="FFFFFF")
FILL_TOTAL = PatternFill(fill_type="solid", fgColor="D0E4F7")
FILL_PROFIT = PatternFill(fill_type="solid", fgColor="E8F5E9")
FILL_SALE = PatternFill(fill_type="solid", fgColor="E3F2FD")
FILL_WARRANTY = PatternFill(fill_type="solid", fgColor="FFF3E0")
FILL_COURTESY = PatternFill(fill_type="solid", fgColor="EDE7F6")
FILL_APPROVED = PatternFill(fill_type="solid", fgColor="E8F5E9")
FILL_ATTENTION = PatternFill(fill_type="solid", fgColor="FFF3E0")
FILL_REJECTED = PatternFill(fill_type="solid", fgColor="FFEBEE")

FONT_TITLE = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
FONT_SUBTITLE = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
FONT_HEADER = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
FONT_BODY = Font(name="Calibri", size=10)
FONT_ID = Font(name="Calibri", size=10, bold=True, color="1565C0")
FONT_SALE = Font(name="Calibri", size=10, bold=True, color="1A3C6E")
FONT_COST = Font(name="Calibri", size=10, color="C62828")
FONT_PROFIT = Font(name="Calibri", size=10, bold=True, color="1B5E20")
FONT_TOTAL = Font(name="Calibri", size=11, bold=True, color="0D2137")
FONT_SECTION_SALE = Font(name="Calibri", size=11, bold=True, color="1565C0")
FONT_SECTION_WARRANTY = Font(name="Calibri", size=11, bold=True, color="E65100")
FONT_SECTION_COURTESY = Font(name="Calibri", size=11, bold=True, color="4527A0")
FONT_BADGE_SALE = Font(name="Calibri", size=10, bold=True, color="1565C0")
FONT_BADGE_WARRANTY = Font(name="Calibri", size=10, bold=True, color="E65100")
FONT_BADGE_COURTESY = Font(name="Calibri", size=10, bold=True, color="4527A0")
FONT_BADGE_APPROVED = Font(name="Calibri", size=10, bold=True, color="2E7D32")
FONT_BADGE_ATTENTION = Font(name="Calibri", size=10, bold=True, color="E65100")
FONT_BADGE_REJECTED = Font(name="Calibri", size=10, bold=True, color="C62828")
FONT_EMPTY = Font(name="Calibri", size=10, italic=True, color="1A3C6E")

THIN_BORDER = Border(
    left=Side(style="thin", color="BFDBFE"),
    right=Side(style="thin", color="BFDBFE"),
    top=Side(style="thin", color="BFDBFE"),
    bottom=Side(style="thin", color="BFDBFE"),
)

CellKind = Literal["text", "id", "date", "money_sale", "money_cost", "money_profit", "percent", "badge"]
BadgeKey = Literal["venda", "garantia", "cortesia", "aprovado", "reprovado", "cancelado", "atencao"]
SectionKey = Literal["venda", "garantia", "cortesia"]


@dataclass(frozen=True, slots=True)
class ExcelColumn:
    header: str
    width: float
    kind: CellKind = "text"


@dataclass(frozen=True, slots=True)
class ExcelCell:
    value: object = None
    kind: CellKind | None = None
    badge: BadgeKey | None = None
    row_fill: Literal["white", "zebra"] | None = None


@dataclass(frozen=True, slots=True)
class ExcelSection:
    title: str
    palette: SectionKey
    rows: list[list[ExcelCell]]
    subtotal_label: str
    subtotal_cells: list[ExcelCell]


def as_excel_number(value: object) -> float:
    amount = getattr(value, "amount", value)
    if isinstance(amount, Decimal):
        return float(amount)
    if isinstance(amount, (int, float)):
        return float(amount)
    if isinstance(amount, str) and amount.strip():
        return float(amount.replace(".", "").replace(",", ".") if "," in amount and amount.count(",") == 1 else amount)
    return 0.0


def percent_ratio(value: object) -> float:
    number = as_excel_number(value)
    if abs(number) > 1:
        return number / 100.0
    return number


def build_hunter_excel_document(
    *,
    title: str,
    workshop_name: str,
    generated_at_label: str,
    count_label: str,
    count: int,
    sheet_title: str,
    columns: list[ExcelColumn],
    rows: list[list[ExcelCell]] | None = None,
    sections: list[ExcelSection] | None = None,
    total_label: str,
    total_cells: list[ExcelCell],
    filename: str,
) -> DocumentPayload:
    workbook = Workbook()
    worksheet = workbook.active
    if worksheet is None:
        raise ValueError("Nao foi possivel inicializar a planilha.")
    worksheet.title = sheet_title[:31]
    _populate_sheet(
        worksheet=worksheet,
        title=title,
        workshop_name=workshop_name,
        generated_at_label=generated_at_label,
        count_label=count_label,
        count=count,
        columns=columns,
        rows=rows or [],
        sections=sections or [],
        total_label=total_label,
        total_cells=total_cells,
    )
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return DocumentPayload(content=buffer.getvalue(), filename=filename, content_type=EXCEL_CONTENT_TYPE)


def _populate_sheet(
    *,
    worksheet: Any,
    title: str,
    workshop_name: str,
    generated_at_label: str,
    count_label: str,
    count: int,
    columns: list[ExcelColumn],
    rows: list[list[ExcelCell]],
    sections: list[ExcelSection],
    total_label: str,
    total_cells: list[ExcelCell],
) -> None:
    column_count = max(1, len(columns))
    last_column = get_column_letter(column_count)
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A4"
    worksheet.row_dimensions[1].height = 38
    worksheet.row_dimensions[2].height = 20
    worksheet.row_dimensions[3].height = 30

    if column_count > 1:
        worksheet.merge_cells(f"A1:{last_column}1")
        worksheet.merge_cells(f"A2:{last_column}2")

    title_cell = worksheet["A1"]
    title_cell.value = title
    title_cell.fill = FILL_TITLE
    title_cell.font = FONT_TITLE
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    subtitle_cell = worksheet["A2"]
    subtitle_cell.value = f"Gerado em: {generated_at_label}   |   {count_label}: {count}   |   {workshop_name}"
    subtitle_cell.fill = FILL_SUBTITLE
    subtitle_cell.font = FONT_SUBTITLE
    subtitle_cell.alignment = Alignment(horizontal="center", vertical="center")

    for column_index, column in enumerate(columns, start=1):
        cell = worksheet.cell(row=3, column=column_index, value=column.header)
        cell.fill = FILL_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
        worksheet.column_dimensions[get_column_letter(column_index)].width = column.width

    worksheet.auto_filter.ref = f"A3:{last_column}3"

    current_row = 4
    if sections:
        for section in sections:
            current_row = _write_section_header(worksheet=worksheet, row=current_row, column_count=column_count, section=section)
            if not section.rows:
                current_row = _write_empty_row(worksheet=worksheet, row=current_row, column_count=column_count)
            else:
                for row_offset, row_cells in enumerate(section.rows):
                    _write_data_row(
                        worksheet=worksheet,
                        row=current_row,
                        columns=columns,
                        cells=row_cells,
                        zebra=row_offset % 2 == 1,
                    )
                    current_row += 1
            current_row = _write_subtotal_row(
                worksheet=worksheet,
                row=current_row,
                columns=columns,
                section=section,
            )
    elif not rows:
        current_row = _write_empty_row(worksheet=worksheet, row=current_row, column_count=column_count)
    else:
        for row_offset, row_cells in enumerate(rows):
            _write_data_row(
                worksheet=worksheet,
                row=current_row,
                columns=columns,
                cells=row_cells,
                zebra=row_offset % 2 == 1,
            )
            current_row += 1

    _write_total_row(
        worksheet=worksheet,
        row=current_row,
        columns=columns,
        total_label=total_label,
        total_cells=total_cells,
        column_count=column_count,
    )


def _write_section_header(*, worksheet: Any, row: int, column_count: int, section: ExcelSection) -> int:
    last_column = get_column_letter(column_count)
    if column_count > 1:
        worksheet.merge_cells(f"A{row}:{last_column}{row}")
    fill, font = _section_style(section.palette)
    cell = worksheet.cell(row=row, column=1, value=section.title)
    cell.fill = fill
    cell.font = font
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.border = THIN_BORDER
    worksheet.row_dimensions[row].height = 26
    return row + 1


def _write_subtotal_row(*, worksheet: Any, row: int, columns: list[ExcelColumn], section: ExcelSection) -> int:
    fill, font = _section_style(section.palette)
    for column_index, column in enumerate(columns, start=1):
        cell_spec = section.subtotal_cells[column_index - 1] if column_index - 1 < len(section.subtotal_cells) else ExcelCell()
        excel_cell = worksheet.cell(row=row, column=column_index)
        if column_index == 1 and cell_spec.value is None:
            excel_cell.value = section.subtotal_label
            excel_cell.font = font
            excel_cell.alignment = Alignment(horizontal="left", vertical="center")
        else:
            _apply_cell(excel_cell=excel_cell, spec=cell_spec, column=column, zebra_fill=fill)
            if cell_spec.kind in {"money_sale", "money_cost", "money_profit", "percent"} or column.kind in {"money_sale", "money_cost", "money_profit", "percent"}:
                excel_cell.fill = fill
        excel_cell.fill = fill
        excel_cell.border = THIN_BORDER
    worksheet.row_dimensions[row].height = 18
    return row + 1


def _write_total_row(
    *,
    worksheet: Any,
    row: int,
    columns: list[ExcelColumn],
    total_label: str,
    total_cells: list[ExcelCell],
    column_count: int,
) -> None:
    del column_count
    worksheet.row_dimensions[row].height = 22
    for column_index, column in enumerate(columns, start=1):
        cell_spec = total_cells[column_index - 1] if column_index - 1 < len(total_cells) else ExcelCell()
        excel_cell = worksheet.cell(row=row, column=column_index)
        if column_index == 1 and cell_spec.value is None:
            excel_cell.value = total_label
            excel_cell.font = FONT_TOTAL
            excel_cell.alignment = Alignment(horizontal="left", vertical="center")
        else:
            _apply_cell(excel_cell=excel_cell, spec=cell_spec, column=column, zebra_fill=FILL_TOTAL)
        excel_cell.fill = FILL_TOTAL
        excel_cell.border = THIN_BORDER
        if column_index == 1:
            excel_cell.font = FONT_TOTAL


def _write_empty_row(*, worksheet: Any, row: int, column_count: int) -> int:
    last_column = get_column_letter(column_count)
    if column_count > 1:
        worksheet.merge_cells(f"A{row}:{last_column}{row}")
    cell = worksheet.cell(row=row, column=1, value="Nenhum registro encontrado para os filtros selecionados.")
    cell.fill = FILL_ZEBRA
    cell.font = FONT_EMPTY
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = THIN_BORDER
    return row + 1


def _write_data_row(*, worksheet: Any, row: int, columns: list[ExcelColumn], cells: list[ExcelCell], zebra: bool) -> None:
    requested_row_fill = next((cell.row_fill for cell in cells if cell.row_fill is not None), None)
    zebra_fill = FILL_ZEBRA if requested_row_fill == "zebra" or (requested_row_fill is None and zebra) else FILL_WHITE
    worksheet.row_dimensions[row].height = 18
    for column_index, column in enumerate(columns, start=1):
        spec = cells[column_index - 1] if column_index - 1 < len(cells) else ExcelCell()
        excel_cell = worksheet.cell(row=row, column=column_index)
        _apply_cell(excel_cell=excel_cell, spec=spec, column=column, zebra_fill=zebra_fill)


def _apply_cell(*, excel_cell: Any, spec: ExcelCell, column: ExcelColumn, zebra_fill: Any) -> None:
    kind = spec.kind or column.kind
    excel_cell.border = THIN_BORDER
    excel_cell.fill = zebra_fill

    if kind == "badge" or spec.badge:
        label, fill, font = _badge_style(spec.badge, spec.value)
        excel_cell.value = label
        excel_cell.fill = fill
        excel_cell.font = font
        excel_cell.alignment = Alignment(horizontal="center", vertical="center")
        return

    if kind == "id":
        excel_cell.value = spec.value
        excel_cell.font = FONT_ID
        excel_cell.alignment = Alignment(horizontal="center", vertical="center")
        return

    if kind == "date":
        excel_cell.value = _as_date(spec.value)
        excel_cell.font = FONT_BODY
        excel_cell.alignment = Alignment(horizontal="center", vertical="center")
        if excel_cell.value:
            excel_cell.number_format = DATE_FORMAT
        return

    if kind == "money_sale":
        excel_cell.value = as_excel_number(spec.value)
        excel_cell.number_format = CURRENCY_FORMAT
        excel_cell.font = FONT_SALE
        excel_cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=False)
        return

    if kind == "money_cost":
        excel_cell.value = as_excel_number(spec.value)
        excel_cell.number_format = CURRENCY_FORMAT
        excel_cell.font = FONT_COST
        excel_cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=False)
        return

    if kind == "money_profit":
        excel_cell.value = as_excel_number(spec.value)
        excel_cell.number_format = CURRENCY_FORMAT
        excel_cell.font = FONT_PROFIT
        excel_cell.fill = FILL_PROFIT
        excel_cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=False)
        return

    if kind == "percent":
        excel_cell.value = percent_ratio(spec.value)
        excel_cell.number_format = PERCENT_FORMAT
        excel_cell.font = FONT_PROFIT
        excel_cell.fill = FILL_PROFIT
        excel_cell.alignment = Alignment(horizontal="right", vertical="center")
        return

    excel_cell.value = "" if spec.value is None else spec.value
    excel_cell.font = FONT_BODY
    excel_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _as_date(value: object) -> date | datetime | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _section_style(palette: SectionKey) -> tuple[Any, Any]:
    if palette == "garantia":
        return FILL_WARRANTY, FONT_SECTION_WARRANTY
    if palette == "cortesia":
        return FILL_COURTESY, FONT_SECTION_COURTESY
    return FILL_SALE, FONT_SECTION_SALE


def _badge_style(badge: BadgeKey | None, value: object) -> tuple[str, Any, Any]:
    label = str(value or "-")
    resolved = badge or _infer_badge(label)
    if resolved == "garantia":
        return label, FILL_WARRANTY, FONT_BADGE_WARRANTY
    if resolved == "cortesia":
        return label, FILL_COURTESY, FONT_BADGE_COURTESY
    if resolved == "aprovado":
        return label, FILL_APPROVED, FONT_BADGE_APPROVED
    if resolved in {"reprovado", "cancelado"}:
        return label, FILL_REJECTED, FONT_BADGE_REJECTED
    if resolved == "atencao":
        return label, FILL_ATTENTION, FONT_BADGE_ATTENTION
    return label, FILL_SALE, FONT_BADGE_SALE


def _infer_badge(label: str) -> BadgeKey:
    normalized = label.strip().casefold()
    if "garantia" in normalized:
        return "garantia"
    if "cortesia" in normalized:
        return "cortesia"
    if normalized in {"aprovado", "aprovada", "veículo entregue", "veiculo entregue"}:
        return "aprovado"
    if "reprov" in normalized:
        return "reprovado"
    if "cancel" in normalized:
        return "cancelado"
    if "aguard" in normalized:
        return "atencao"
    return "venda"
