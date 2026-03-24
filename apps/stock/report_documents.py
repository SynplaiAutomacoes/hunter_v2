from __future__ import annotations

from io import BytesIO
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.core.documents.contract import DocumentPayload, DocumentRenderRequest
from apps.core.documents.renderer import render_template_request_to_pdf
from apps.stock.reporting import StockReportColumnDefinition
from apps.workshops.models.workshops import Workshop


_STOCK_REPORT_EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_PRIMARY_FILL = PatternFill(fill_type="solid", fgColor="1F2937")
_SECONDARY_FILL = PatternFill(fill_type="solid", fgColor="F9FAFB")
_HIGHLIGHT_FILL = PatternFill(fill_type="solid", fgColor="EFF6FF")
_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)


def build_stock_report_pdf_render_request(*, context: dict[str, object], request=None, filename: str | None = None) -> DocumentRenderRequest:
    workshop = context.get("workshop") if isinstance(context.get("workshop"), Workshop) else None
    resolved_filename = filename or _build_filename(workshop=workshop, extension="pdf")
    render_context = dict(context)
    render_context["request"] = request

    return DocumentRenderRequest(
        template_name="stock/report_pdf.html",
        context=render_context,
        filename=resolved_filename,
    )


def render_stock_report_pdf_document(*, context: dict[str, object], request=None, filename: str | None = None) -> DocumentPayload:
    render_request = build_stock_report_pdf_render_request(context=context, request=request, filename=filename)
    return render_template_request_to_pdf(render_request)


def build_stock_report_excel_document(*, context: dict[str, object], filename: str | None = None) -> DocumentPayload:
    workbook = Workbook()
    worksheet = workbook.active
    if worksheet is None:
        raise ValueError("Nao foi possivel inicializar a planilha do relatorio de estoque.")

    worksheet.title = "Relatorio"
    _populate_stock_report_sheet(worksheet=worksheet, context=context)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    workshop = context.get("workshop") if isinstance(context.get("workshop"), Workshop) else None
    resolved_filename = filename or _build_filename(workshop=workshop, extension="xlsx")
    return DocumentPayload(content=buffer.getvalue(), filename=resolved_filename, content_type=_STOCK_REPORT_EXCEL_CONTENT_TYPE)


def _populate_stock_report_sheet(*, worksheet, context: dict[str, object]) -> None:
    selected_columns = [column for column in context.get("selected_columns", []) if isinstance(column, StockReportColumnDefinition)]
    stock_report_items = [item for item in context.get("stock_report_items", []) if hasattr(item, "product")]
    stock_report_totals = context.get("stock_report_totals") if isinstance(context.get("stock_report_totals"), dict) else {}
    filter_descriptions = [str(item) for item in context.get("stock_report_filter_descriptions", []) if str(item).strip()]
    workshop = context.get("workshop") if isinstance(context.get("workshop"), Workshop) else None
    generated_at_label = str(context.get("generated_at_label") or "-")

    column_count = max(1, len(selected_columns))
    last_column_letter = get_column_letter(column_count)
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A10"

    if column_count > 1:
        worksheet.merge_cells(f"A1:{last_column_letter}1")
    title_cell = worksheet["A1"]
    title_cell.value = "Relatorio de Estoque"
    title_cell.fill = _PRIMARY_FILL
    title_cell.font = Font(color="FFFFFF", bold=True, size=16)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 28

    metadata_rows = [
        ("Oficina", workshop.name if workshop is not None else "-"),
        ("Emitido em", generated_at_label),
        ("Filtros", ", ".join(filter_descriptions) if filter_descriptions else "Sem filtros adicionais"),
        ("Itens filtrados", stock_report_totals.get("item_count_display", "0")),
        ("Quantidade total", stock_report_totals.get("total_quantity_display", "0")),
        ("Custo total do estoque", stock_report_totals.get("stock_total_cost_display", "R$ 0,00")),
    ]

    for row_index, (label, value) in enumerate(metadata_rows, start=3):
        worksheet[f"A{row_index}"] = label
        worksheet[f"B{row_index}"] = value
        worksheet[f"A{row_index}"].font = Font(bold=True, color="1F2937")
        worksheet[f"A{row_index}"].fill = _SECONDARY_FILL
        worksheet[f"A{row_index}"].border = _BORDER
        worksheet[f"B{row_index}"].border = _BORDER
        worksheet[f"B{row_index}"].alignment = Alignment(wrap_text=True)

    header_row = 10
    for column_index, column in enumerate(selected_columns, start=1):
        cell = worksheet.cell(row=header_row, column=column_index, value=column.label)
        cell.fill = _PRIMARY_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = _BORDER

    if not stock_report_items:
        if column_count > 1:
            worksheet.merge_cells(f"A11:{last_column_letter}11")
        empty_cell = worksheet["A11"]
        empty_cell.value = "Nenhum item encontrado para os filtros informados."
        empty_cell.fill = _SECONDARY_FILL
        empty_cell.alignment = Alignment(horizontal="center")
        empty_cell.font = Font(italic=True, color="374151")
        empty_cell.border = _BORDER
    else:
        for row_index, item in enumerate(stock_report_items, start=header_row + 1):
            for column_index, column in enumerate(selected_columns, start=1):
                cell = worksheet.cell(row=row_index, column=column_index, value=column.excel_value_resolver(item))
                cell.border = _BORDER
                cell.fill = _HIGHLIGHT_FILL if row_index % 2 == 0 else _SECONDARY_FILL
                cell.alignment = Alignment(horizontal="right" if column.pdf_align == "right" else "left", vertical="center")
                if column.excel_number_format:
                    cell.number_format = column.excel_number_format

    for column_index, column in enumerate(selected_columns, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = column.excel_width

    worksheet.column_dimensions["A"].width = max(worksheet.column_dimensions["A"].width, 18)
    worksheet.column_dimensions["B"].width = max(worksheet.column_dimensions["B"].width, 28)


def _build_filename(*, workshop: Workshop | None, extension: str) -> str:
    workshop_fragment = _normalize_filename_fragment(workshop.name) if workshop is not None else "oficina"
    return f"relatorio_estoque_{workshop_fragment}.{extension}"


def _normalize_filename_fragment(value: str) -> str:
    normalized_value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    normalized_value = normalized_value.lower().strip()
    normalized_value = re.sub(r"[^a-z0-9]+", "_", normalized_value)
    normalized_value = normalized_value.strip("_")
    return normalized_value or "-"
