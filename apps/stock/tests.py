from __future__ import annotations

import base64
import gzip
from datetime import timedelta
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from django.core.management import call_command
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import RequestFactory, TestCase
from django.urls import reverse
from django.utils import timezone
from django.template import Context, Template
from openpyxl import load_workbook
from djmoney.money import Money

from apps.accounts.models import Account, User
from apps.catalog.models.groups import CatalogGroup
from apps.catalog.models.products import Product
from apps.collaborators.models import WorkshopMember
from apps.core.documents.contract import DocumentPayload
from apps.finance.models.payment_method import PaymentMethod
from apps.iam.utils import get_or_create_director_role
from apps.stock.forms import ImportManualItemsForm, ImportSefazListForm, ImportStep1Form, ImportStepPaymentForm, ImportStepSummaryForm, ImportStepSupplierForm, QuickProductForm
from apps.stock.models import SefazZipCache, StockImport, StockMovement, StockPaymentMethod, StockProduct, StockTransfer
from apps.stock.utils import NFParser
from apps.suppliers.models import Supplier
from apps.workshops.models.workshops import Workshop


class StockSefazTests(TestCase):
    def setUp(self) -> None:
        self.factory = RequestFactory()
        self.workshop = Workshop.objects.create(
            name="Oficina Teste",
            cnpj="11.222.333/0001-81",
            phone="+5511999999999",
            address="Rua Teste, 123",
            uf="SP",
            pfx_certificate=SimpleUploadedFile("teste.pfx", b"certificado-teste", content_type="application/x-pkcs12"),
            certificate_password="segredo",
        )

    @staticmethod
    def _build_access_key(*, nf_number: int) -> str:
        return f"3526031122233300018155001{nf_number:09d}1123456789"

    def _build_res_nfe_xml(self, *, nf_number: int = 123, include_number: bool = True) -> bytes:
        access_key = self._build_access_key(nf_number=nf_number)
        number_xml = f"<nNF>{nf_number}</nNF>" if include_number else ""
        return f"""<?xml version=\"1.0\" encoding=\"UTF-8\"?>
<resNFe xmlns=\"http://www.portalfiscal.inf.br/nfe\">
    <chNFe>{access_key}</chNFe>
    {number_xml}
    <xNome>Fornecedor Teste</xNome>
    <CNPJ>11222333000181</CNPJ>
    <vNF>123.45</vNF>
    <dhEmi>2026-03-20T10:00:00-03:00</dhEmi>
</resNFe>
""".encode()

    def _build_distribution_response(self, *, document_xml: bytes) -> bytes:
        encoded_doc = base64.b64encode(gzip.compress(document_xml)).decode()
        return f"""<?xml version=\"1.0\" encoding=\"UTF-8\"?>
<retDistDFeInt xmlns=\"http://www.portalfiscal.inf.br/nfe\">
    <cStat>138</cStat>
    <xMotivo>Documentos localizados</xMotivo>
    <ultNSU>000000000000001</ultNSU>
    <maxNSU>000000000000001</maxNSU>
    <loteDistDFeInt>
        <docZip NSU=\"000000000000001\" schema=\"resNFe_v1.01.xsd\">{encoded_doc}</docZip>
    </loteDistDFeInt>
</retDistDFeInt>
""".encode()

    def test_parse_res_nfe_reads_number_from_element(self) -> None:
        parsed = NFParser.parse_nfe_xml_to_dict(self.workshop, self._build_res_nfe_xml())

        self.assertIsNotNone(parsed)
        assert parsed is not None
        self.assertEqual(parsed["nf_number"], "123")
        self.assertEqual(parsed["nf_key"], self._build_access_key(nf_number=123))

    def test_parse_res_nfe_falls_back_to_number_from_access_key(self) -> None:
        parsed = NFParser.parse_nfe_xml_to_dict(self.workshop, self._build_res_nfe_xml(include_number=False))

        self.assertIsNotNone(parsed)
        assert parsed is not None
        self.assertEqual(parsed["nf_number"], "123")

    def test_update_sefaz_list_persists_number_in_cache(self) -> None:
        response_content = self._build_distribution_response(document_xml=self._build_res_nfe_xml(include_number=False))
        form = ImportSefazListForm(workshop=self.workshop)

        with patch("apps.stock.forms.ComunicacaoSefaz") as comunicacao_cls:
            comunicacao_cls.return_value.consulta_distribuicao.return_value = SimpleNamespace(content=response_content)

            success, _message = form.update_sefaz_list()

        self.assertTrue(success)
        cache = SefazZipCache.objects.get(workshop=self.workshop, key=self._build_access_key(nf_number=123))
        self.assertEqual(cache.nf_number, "123")
        self.assertEqual(cache.nf_number_display, "123")

    def test_import_step1_save_uses_access_key_fallback_for_nf_number(self) -> None:
        access_key = self._build_access_key(nf_number=321)
        form = ImportStep1Form(workshop=self.workshop, data={"method": "KEY"})
        form.cleaned_data = {"method": "KEY"}
        form.parsed_nf_data = {
            "nf_key": access_key,
            "nf_number": None,
            "supplier_cnpj": "11222333000181",
            "supplier_name": "Fornecedor Teste",
            "items": [],
            "payments": [],
        }

        stock_import = form.save(commit=False)

        self.assertEqual(stock_import.nf_number, "321")

    def test_import_sefaz_list_form_uses_paginated_page_queryset(self) -> None:
        issued_at = timezone.now()
        for nf_number in range(1, 16):
            SefazZipCache.objects.create(
                workshop=self.workshop,
                key=self._build_access_key(nf_number=nf_number),
                nf_number=str(nf_number),
                issuer_name=f"Fornecedor {nf_number}",
                issuer_cnpj="11222333000181",
                total_value="123.45",
                issue_date=issued_at + timedelta(minutes=nf_number),
            )

        request = self.factory.get("/stock/import/", {"page": 2})
        form = ImportSefazListForm(workshop=self.workshop, request=request)

        self.assertEqual(form.notas.number, 2)
        self.assertEqual(len(form.notas.object_list), 5)
        self.assertEqual([nota.nf_number for nota in form.notas.object_list], ["5", "4", "3", "2", "1"])

    def test_import_sefaz_save_updates_cache_with_resolved_number(self) -> None:
        access_key = self._build_access_key(nf_number=456)
        cache = SefazZipCache.objects.create(workshop=self.workshop, key=access_key)
        form = ImportSefazListForm(workshop=self.workshop, data={"selected_key": access_key})

        self.assertTrue(form.is_valid())

        with patch("apps.stock.forms.ComunicacaoSefaz") as comunicacao_cls:
            comunicacao_cls.return_value.consulta_distribuicao.return_value = SimpleNamespace(content=self._build_res_nfe_xml(nf_number=456, include_number=False))

            stock_import = form.save(commit=False)

        cache.refresh_from_db()
        self.assertEqual(stock_import.nf_number, "456")
        self.assertEqual(cache.nf_number, "456")
        self.assertEqual(cache.issuer_name, "Fornecedor Teste")
        self.assertEqual(cache.issuer_cnpj, "11222333000181")

    def test_supplier_and_summary_forms_use_nf_number_display_fallback(self) -> None:
        stock_import = StockImport.objects.create(
            workshop=self.workshop,
            nf_key=self._build_access_key(nf_number=654),
            supplier_name="Fornecedor Teste",
            supplier_cnpj="11222333000181",
            items_data=[],
            payments_data=[],
        )

        supplier_form = ImportStepSupplierForm(instance=stock_import, workshop=self.workshop)
        summary_form = ImportStepSummaryForm(instance=stock_import, workshop=self.workshop)

        assert supplier_form.helper is not None
        assert supplier_form.helper.layout is not None
        assert summary_form.helper is not None
        assert summary_form.helper.layout is not None
        assert supplier_form.helper.layout.fields
        assert summary_form.helper.layout.fields

        self.assertIn("NF-e: 654", supplier_form.helper.layout.fields[0].html)
        self.assertIn("#654", summary_form.helper.layout.fields[0].html)

    def test_stock_import_number_display_falls_back_to_access_key(self) -> None:
        stock_import = StockImport.objects.create(
            workshop=self.workshop,
            nf_key=self._build_access_key(nf_number=987),
        )

        self.assertEqual(stock_import.nf_number_display, "987")


def create_director_user_with_workshop(*, suffix: int = 1) -> tuple[User, Workshop]:
    user = User.objects.create_user(username=f"stock_director{suffix}", password="123", cpf=f"33344455{suffix:03d}")
    account = Account.objects.create(name=f"Conta Estoque {suffix}", owner=user)
    user.account = account
    user.is_account_owner = True
    user.save(update_fields=["account", "is_account_owner"])

    workshop = Workshop.objects.create(
        account=account,
        name=f"Oficina Estoque {suffix}",
        cnpj=f"31.222.444/0001-{suffix:02d}",
        phone="+5511988887777",
        address="Rua Estoque, 10",
        uf="SP",
    )

    director_role = get_or_create_director_role(account=account, with_all_permissions=True)
    WorkshopMember.objects.create(user=user, workshop=workshop, role=director_role, is_active=True)
    return user, workshop


class StockTransferFlowTests(TestCase):
    def setUp(self) -> None:
        self.user, self.destination_workshop = create_director_user_with_workshop(suffix=20)
        self.source_workshop = Workshop.objects.create(
            account=self.destination_workshop.account,
            name="Oficina Origem",
            cnpj="31.222.444/0001-21",
            phone="+5511988887766",
            address="Rua Origem, 20",
            uf="SP",
        )
        director_role = get_or_create_director_role(account=self.destination_workshop.account, with_all_permissions=True)
        WorkshopMember.objects.create(user=self.user, workshop=self.source_workshop, role=director_role, is_active=True)

        self.source_group = CatalogGroup.objects.create(workshop=self.source_workshop, name="Filtros")
        self.destination_group = CatalogGroup.objects.create(workshop=self.destination_workshop, name="Filtros")
        self.source_product = Product.objects.create(
            workshop=self.source_workshop,
            code="FLT-001",
            name="Filtro de Oleo",
            unit=Product.Unit.UND,
            group=self.source_group,
            cost_price="10.00",
            selling_price="15.00",
            origin_cst=Product.OriginCST.NACIONAL,
            purpose=Product.Purpose.RESALE,
        )
        self.destination_product = Product.objects.create(
            workshop=self.destination_workshop,
            code="FLT-DEST-001",
            name="Filtro de Oleo Destino",
            unit=Product.Unit.UND,
            group=self.destination_group,
            cost_price="11.00",
            selling_price="16.00",
            origin_cst=Product.OriginCST.NACIONAL,
            purpose=Product.Purpose.RESALE,
        )
        self.source_stock = StockProduct.objects.get(workshop=self.source_workshop, product=self.source_product)
        self.source_stock.current_quantity = 8
        self.source_stock.save(update_fields=["current_quantity"])
        self.destination_stock = StockProduct.objects.get(workshop=self.destination_workshop, product=self.destination_product)
        self.destination_stock.current_quantity = 2
        self.destination_stock.save(update_fields=["current_quantity"])

        self.client.force_login(self.user)
        session = self.client.session
        session["active_workshop_id"] = self.destination_workshop.pk
        session.save()

    def test_transfer_workflow_moves_stock_between_workshops(self) -> None:
        response = self.client.post(reverse("stock:transfer") + "?step=1", {"source_workshop": self.source_workshop.pk, "destination_workshop": self.destination_workshop.pk})
        self.assertEqual(response.status_code, 302)

        transfer = StockTransfer.objects.get()
        transfer.items_data = [
            {
                "source_product_id": str(self.source_product.pk),
                "destination_product_id": str(self.destination_product.pk),
                "qtd": 3,
                "valor": "10.00",
            }
        ]
        transfer.current_step = 3
        transfer.save(update_fields=["items_data", "current_step"])

        response = self.client.post(reverse("stock:transfer_update", kwargs={"pk": transfer.pk}) + "?step=3", {})
        self.assertEqual(response.status_code, 302)

        transfer.refresh_from_db()
        self.source_stock.refresh_from_db()
        self.destination_stock.refresh_from_db()

        self.assertEqual(transfer.status, StockTransfer.TransferStatus.COMPLETED)
        self.assertEqual(self.source_stock.current_quantity, 5)
        self.assertEqual(self.destination_stock.current_quantity, 5)

        exit_movement = StockMovement.objects.get(workshop=self.source_workshop, stock_transfer=transfer, type=StockMovement.MovementType.EXIT)
        entry_movement = StockMovement.objects.get(workshop=self.destination_workshop, stock_transfer=transfer, type=StockMovement.MovementType.ENTRY)
        self.assertEqual(exit_movement.quantity, 3)
        self.assertEqual(entry_movement.quantity, 3)

    def test_transfer_summary_blocks_when_source_stock_is_insufficient(self) -> None:
        transfer = StockTransfer.objects.create(
            source_workshop=self.source_workshop,
            destination_workshop=self.destination_workshop,
            user=self.user,
            current_step=3,
            items_data=[
                {
                    "source_product_id": str(self.source_product.pk),
                    "destination_product_id": str(self.destination_product.pk),
                    "qtd": 99,
                    "valor": "10.00",
                }
            ],
        )

        response = self.client.post(reverse("stock:transfer_update", kwargs={"pk": transfer.pk}) + "?step=3", {})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Saldo insuficiente")

        transfer.refresh_from_db()
        self.source_stock.refresh_from_db()
        self.destination_stock.refresh_from_db()
        self.assertEqual(transfer.status, StockTransfer.TransferStatus.DRAFT)
        self.assertEqual(self.source_stock.current_quantity, 8)
        self.assertEqual(self.destination_stock.current_quantity, 2)

    def test_create_destination_product_endpoint_clones_product_into_destination(self) -> None:
        transfer = StockTransfer.objects.create(
            source_workshop=self.source_workshop,
            destination_workshop=self.destination_workshop,
            user=self.user,
            current_step=2,
            items_data=[
                {
                    "source_product_id": str(self.source_product.pk),
                    "destination_product_id": None,
                    "qtd": 1,
                    "valor": "10.00",
                }
            ],
        )

        response = self.client.post(reverse("stock:create_transfer_destination_product") + f"?pk={transfer.pk}&item_idx=0")
        self.assertEqual(response.status_code, 204)

        transfer.refresh_from_db()
        destination_product_id = transfer.items_data[0]["destination_product_id"]
        self.assertIsNotNone(destination_product_id)

        created_product = Product.objects.get(pk=destination_product_id)
        self.assertEqual(created_product.workshop, self.destination_workshop)
        self.assertEqual(created_product.code, self.source_product.code)
        self.assertEqual(created_product.group.name, self.source_group.name)

    def test_stock_history_lists_imports_and_transfers(self) -> None:
        stock_import = StockImport.objects.create(
            workshop=self.destination_workshop,
            user=self.user,
            nf_key="0" * 44,
            nf_number="1234",
            supplier_name="Fornecedor XPTO",
            status=StockImport.ImportStatus.COMPLETED,
        )
        transfer = StockTransfer.objects.create(
            source_workshop=self.source_workshop,
            destination_workshop=self.destination_workshop,
            user=self.user,
            status=StockTransfer.TransferStatus.COMPLETED,
        )

        response = self.client.get(reverse("stock:stock_list"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "1234")
        self.assertContains(response, "Fornecedor XPTO")
        self.assertContains(response, "TRANSFERENCIA")
        self.assertContains(response, f"{self.source_workshop.name} -&gt; {self.destination_workshop.name}")
        self.assertContains(response, reverse("stock:history_edit", kwargs={"record_type": "import", "pk": stock_import.pk}))
        self.assertContains(response, reverse("stock:history_edit", kwargs={"record_type": "transfer", "pk": transfer.pk}))
        self.assertContains(response, reverse("stock:stock_delete", args=[stock_import.pk]))
        self.assertEqual(response.content.decode("utf-8").count("btn-table-delete"), 2)

    def test_history_edit_redirect_routes_transfer_to_transfer_update(self) -> None:
        transfer = StockTransfer.objects.create(
            source_workshop=self.source_workshop,
            destination_workshop=self.destination_workshop,
            user=self.user,
        )

        response = self.client.get(reverse("stock:history_edit", kwargs={"record_type": "transfer", "pk": transfer.pk}))

        self.assertRedirects(response, reverse("stock:transfer_update", kwargs={"pk": transfer.pk}), fetch_redirect_response=False)


class QuickProductFormTests(TestCase):
    def setUp(self) -> None:
        self.workshop = Workshop.objects.create(
            name="Oficina Quick Produto",
            cnpj="31.222.555/0001-10",
            phone="+5511988886666",
            address="Rua Quick, 10",
            uf="SP",
        )
        self.group = CatalogGroup.objects.create(workshop=self.workshop, name="Grupo Quick Produto")

    def test_quick_product_form_accepts_optional_ncm(self) -> None:
        form = QuickProductForm(
            data={
                "code": "QP-001",
                "name": "Produto Quick com NCM",
                "unit": Product.Unit.UND,
                "group": str(self.group.pk),
                "cost_price_0": "10.00",
                "cost_price_1": "BRL",
                "selling_price_0": "20.00",
                "selling_price_1": "BRL",
                "profit_margin": "50.00",
                "ncm": "87089990",
                "origin_cst": str(Product.OriginCST.NACIONAL),
                "purpose": Product.Purpose.RESALE,
            },
            workshop=self.workshop,
        )

        self.assertTrue(form.is_valid(), form.errors.as_json())
        product = form.save(commit=False)
        self.assertEqual(product.ncm, "87089990")

    def test_quick_product_form_keeps_ncm_optional(self) -> None:
        form = QuickProductForm(
            data={
                "code": "QP-002",
                "name": "Produto Quick sem NCM",
                "unit": Product.Unit.UND,
                "group": str(self.group.pk),
                "cost_price_0": "10.00",
                "cost_price_1": "BRL",
                "selling_price_0": "20.00",
                "selling_price_1": "BRL",
                "profit_margin": "50.00",
                "ncm": "",
                "origin_cst": str(Product.OriginCST.NACIONAL),
                "purpose": Product.Purpose.RESALE,
            },
            workshop=self.workshop,
        )

        self.assertTrue(form.is_valid(), form.errors.as_json())


class ManualStockImportPricingTests(TestCase):
    def setUp(self) -> None:
        self.user, self.workshop = create_director_user_with_workshop(suffix=31)
        self.group = CatalogGroup.objects.create(workshop=self.workshop, name="Frutas")
        self.product = Product.objects.create(
            workshop=self.workshop,
            code="BAN-001",
            name="Banana",
            unit=Product.Unit.UND,
            group=self.group,
            cost_price=Money("7.00", "BRL"),
            selling_price=Money("12.00", "BRL"),
            last_purchase_price=Money("6.50", "BRL"),
            last_used_price=Money("15.00", "BRL"),
            origin_cst=Product.OriginCST.NACIONAL,
            purpose=Product.Purpose.RESALE,
        )
        self.stock_product = StockProduct.objects.get(workshop=self.workshop, product=self.product)
        self.stock_product.current_quantity = 4
        self.stock_product.save(update_fields=["current_quantity"])
        self.stock_import = StockImport.objects.create(
            workshop=self.workshop,
            user=self.user,
            method=StockImport.ImportMethods.MANUAL,
            current_step=3,
            nf_key="3" * 44,
            supplier_name="Fornecedor Banana",
            supplier_cnpj="12.345.678/0001-31",
            items_data=[
                {
                    "ref": self.product.code,
                    "desc": self.product.name,
                    "qtd": "2",
                    "valor": "7.00",
                    "selling_price": "12.00",
                    "linked_product_id": str(self.product.pk),
                }
            ],
            payments_data=[],
        )

        self.client.force_login(self.user)
        session = self.client.session
        session["active_workshop_id"] = self.workshop.pk
        session.save()

    def test_manual_items_form_renders_selling_and_last_price_columns(self) -> None:
        form = ImportManualItemsForm(instance=self.stock_import, workshop=self.workshop, request=SimpleNamespace(user=self.user))

        html = Template("{% load crispy_forms_tags %}{% crispy form %}").render(Context({"form": form}))
        html = html.replace("\xa0", " ")

        self.assertIn("Valor de Venda", html)
        self.assertIn("Último valor de compra", html)
        self.assertIn("Último valor de venda", html)
        self.assertIn("R$ 6,50", html)
        self.assertIn("R$ 15,00", html)
        self.assertIn("manual-confirm-lower-price-input", html)
        self.assertNotIn('name="items_qty_0"', html)
        self.assertNotIn('name="items_price_0_0"', html)
        self.assertNotIn('name="items_selling_price_0_0"', html)
        self.assertNotIn(reverse("stock:update_manual_item_data", kwargs={"pk": self.stock_import.pk}), html)

    def test_manual_items_form_renders_edit_action_for_existing_row(self) -> None:
        form = ImportManualItemsForm(instance=self.stock_import, workshop=self.workshop, request=SimpleNamespace(user=self.user))

        html = Template("{% load crispy_forms_tags %}{% crispy form %}").render(Context({"form": form}))

        self.assertIn(reverse("stock:manual_link_item_editor"), html)
        self.assertIn(f"product_id={self.product.pk}&item_idx=0", html)
        self.assertIn("Editar Item", html)

    def test_update_manual_item_data_persists_selling_price(self) -> None:
        response = self.client.post(
            reverse("stock:update_manual_item_data", kwargs={"pk": self.stock_import.pk}),
            data={
                "item_idx": "0",
                "items_selling_price_0_0": "18.50",
            },
        )

        self.stock_import.refresh_from_db()

        self.assertEqual(response.status_code, 204)
        self.assertEqual(self.stock_import.items_data[0]["selling_price"], "18.50")

    def test_manual_items_form_requires_confirmation_for_price_below_last_used_price(self) -> None:
        form = ImportManualItemsForm(
            data={
                "items_qty_0": "2",
                "items_price_0_0": "7.00",
                "items_price_0_1": "BRL",
                "items_selling_price_0_0": "10.00",
                "items_selling_price_0_1": "BRL",
            },
            instance=self.stock_import,
            workshop=self.workshop,
            request=SimpleNamespace(user=self.user),
        )

        self.assertFalse(form.is_valid())
        self.assertIn("Banana: Último valor usado: R$ 15,00", str(form.non_field_errors()))
        self.assertEqual(form.instance.items_data[0]["selling_price"], "10.00")

    def test_manual_items_form_allows_confirmed_price_below_last_used_price(self) -> None:
        form = ImportManualItemsForm(
            data={
                "confirm_lower_price": "1",
                "items_qty_0": "2",
                "items_price_0_0": "7.00",
                "items_price_0_1": "BRL",
                "items_selling_price_0_0": "10.00",
                "items_selling_price_0_1": "BRL",
            },
            instance=self.stock_import,
            workshop=self.workshop,
            request=SimpleNamespace(user=self.user),
        )

        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.instance.items_data[0]["selling_price"], "10.00")

    def test_summary_save_updates_product_last_purchase_and_last_used_prices(self) -> None:
        self.product.last_purchase_price = Money("6.50", "BRL")
        self.product.last_used_price = Money("15.00", "BRL")
        self.product.save(update_fields=["last_purchase_price", "last_purchase_price_currency", "last_used_price", "last_used_price_currency"])

        stock_import = StockImport.objects.create(
            workshop=self.workshop,
            user=self.user,
            method=StockImport.ImportMethods.MANUAL,
            nf_key="4" * 44,
            nf_number="NF-BANANA",
            supplier_name="Fornecedor Banana",
            supplier_cnpj="12.345.678/0001-31",
            items_data=[
                {
                    "ref": self.product.code,
                    "desc": self.product.name,
                    "qtd": "3",
                    "valor": "8.00",
                    "selling_price": "13.50",
                    "linked_product_id": str(self.product.pk),
                }
            ],
            payments_data=[],
        )
        form = ImportStepSummaryForm(data={}, instance=stock_import, workshop=self.workshop, request=SimpleNamespace(user=self.user))

        self.assertTrue(form.is_valid(), form.errors)
        form.save()

        self.product.refresh_from_db()
        self.stock_product.refresh_from_db()
        stock_import.refresh_from_db()

        self.assertEqual(stock_import.status, StockImport.ImportStatus.COMPLETED)
        self.assertEqual(self.product.last_purchase_price, Money("8.00", "BRL"))
        self.assertEqual(self.product.last_used_price, Money("13.50", "BRL"))
        self.assertEqual(self.stock_product.current_quantity, 7)


class ManualStockImportLinkEditorTests(TestCase):
    def setUp(self) -> None:
        self.user, self.workshop = create_director_user_with_workshop(suffix=32)
        self.group = CatalogGroup.objects.create(workshop=self.workshop, name="Fila Manual")
        self.product = Product.objects.create(
            workshop=self.workshop,
            code="BAN-FILA",
            name="Banana Fila",
            unit=Product.Unit.UND,
            group=self.group,
            cost_price=Money("8.00", "BRL"),
            selling_price=Money("13.00", "BRL"),
            last_purchase_price=Money("7.25", "BRL"),
            last_used_price=Money("15.00", "BRL"),
            origin_cst=Product.OriginCST.NACIONAL,
            purpose=Product.Purpose.RESALE,
        )
        self.stock_product = StockProduct.objects.get(workshop=self.workshop, product=self.product)
        self.stock_product.current_quantity = 9
        self.stock_product.save(update_fields=["current_quantity"])
        self.stock_import = StockImport.objects.create(
            workshop=self.workshop,
            user=self.user,
            method=StockImport.ImportMethods.MANUAL,
            current_step=3,
            nf_key="5" * 44,
            supplier_name="Fornecedor Fila",
            supplier_cnpj="12.345.678/0001-32",
            items_data=[],
            payments_data=[],
        )

        self.client.force_login(self.user)
        session = self.client.session
        session["active_workshop_id"] = self.workshop.pk
        session.save()

    def test_link_manual_modal_manual_mode_opens_child_editor(self) -> None:
        response = self.client.get(reverse("stock:link_product_manual"), {"pk": self.stock_import.pk, "manual": "true"})

        self.assertContains(response, reverse("stock:manual_link_item_editor"))
        self.assertContains(response, "#child-modal-container")

    def test_manual_link_item_editor_get_prefills_product_information(self) -> None:
        response = self.client.get(
            reverse("stock:manual_link_item_editor"),
            {"pk": self.stock_import.pk, "product_id": self.product.pk},
        )

        self.assertContains(response, "Banana Fila")
        self.assertContains(response, "Último valor de compra")
        self.assertContains(response, "Último valor de venda")
        self.assertContains(response, 'name="quantity"')
        self.assertContains(response, 'name="unit_cost_0"')
        self.assertContains(response, 'name="selling_price_0"')
        self.assertContains(response, 'hx-trigger="manual-link-editor-submit"')
        self.assertContains(response, '@submit.prevent="handleSubmit($event)"')
        self.assertContains(response, "submitForm()")

    def test_manual_link_item_editor_get_with_item_idx_prefills_existing_row(self) -> None:
        self.stock_import.items_data = [
            {
                "ref": self.product.code,
                "desc": self.product.name,
                "qtd": "2",
                "valor": "8.50",
                "selling_price": "16.25",
                "linked_product_id": str(self.product.pk),
            }
        ]
        self.stock_import.save(update_fields=["items_data"])

        response = self.client.get(
            reverse("stock:manual_link_item_editor"),
            {"pk": self.stock_import.pk, "product_id": self.product.pk, "item_idx": "0"},
        )

        self.assertContains(response, "Editar item vinculado")
        self.assertContains(response, 'name="item_idx"')
        self.assertContains(response, 'value="0"')
        self.assertContains(response, 'value="2"')
        self.assertContains(response, 'value="8.50"')
        self.assertContains(response, 'value="16.25"')

    def test_manual_link_item_editor_post_adds_item_after_save(self) -> None:
        response = self.client.post(
            f"{reverse('stock:manual_link_item_editor')}?pk={self.stock_import.pk}&product_id={self.product.pk}",
            data={
                "product_id": str(self.product.pk),
                "quantity": "2",
                "unit_cost_0": "8.50",
                "unit_cost_1": "BRL",
                "selling_price_0": "16.25",
                "selling_price_1": "BRL",
                "confirm_lower_price": "",
            },
        )

        self.stock_import.refresh_from_db()

        self.assertEqual(response.status_code, 200)
        self.assertJSONEqual(
            response.headers["HX-Trigger"],
            {
                "productCreated": {},
                "showToast": {
                    "type": "success",
                    "message": "Banana Fila adicionado a importacao manual.",
                },
            },
        )
        self.assertEqual(len(self.stock_import.items_data), 1)
        self.assertEqual(self.stock_import.items_data[0]["linked_product_id"], str(self.product.pk))
        self.assertEqual(self.stock_import.items_data[0]["qtd"], "2")
        self.assertEqual(self.stock_import.items_data[0]["valor"], "8.50")
        self.assertEqual(self.stock_import.items_data[0]["selling_price"], "16.25")

    def test_manual_link_item_editor_requires_confirmation_for_lower_price(self) -> None:
        response = self.client.post(
            f"{reverse('stock:manual_link_item_editor')}?pk={self.stock_import.pk}&product_id={self.product.pk}",
            data={
                "product_id": str(self.product.pk),
                "quantity": "1",
                "unit_cost_0": "8.00",
                "unit_cost_1": "BRL",
                "selling_price_0": "10.00",
                "selling_price_1": "BRL",
                "confirm_lower_price": "",
            },
        )

        self.stock_import.refresh_from_db()
        response_text = response.content.decode("utf-8").replace("\xa0", " ")

        self.assertIn("Último valor usado: R$ 15,00", response_text)
        self.assertEqual(self.stock_import.items_data, [])

    def test_manual_link_item_editor_allows_confirmed_lower_price(self) -> None:
        response = self.client.post(
            f"{reverse('stock:manual_link_item_editor')}?pk={self.stock_import.pk}&product_id={self.product.pk}",
            data={
                "product_id": str(self.product.pk),
                "quantity": "1",
                "unit_cost_0": "8.00",
                "unit_cost_1": "BRL",
                "selling_price_0": "10.00",
                "selling_price_1": "BRL",
                "confirm_lower_price": "1",
            },
        )

        self.stock_import.refresh_from_db()

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(self.stock_import.items_data), 1)
        self.assertEqual(self.stock_import.items_data[0]["selling_price"], "10.00")

    def test_manual_link_item_editor_post_with_item_idx_updates_existing_row(self) -> None:
        self.stock_import.items_data = [
            {
                "ref": self.product.code,
                "desc": self.product.name,
                "qtd": "1",
                "valor": "8.00",
                "selling_price": "13.00",
                "linked_product_id": str(self.product.pk),
            }
        ]
        self.stock_import.save(update_fields=["items_data"])

        response = self.client.post(
            f"{reverse('stock:manual_link_item_editor')}?pk={self.stock_import.pk}&product_id={self.product.pk}&item_idx=0",
            data={
                "product_id": str(self.product.pk),
                "item_idx": "0",
                "quantity": "3",
                "unit_cost_0": "9.10",
                "unit_cost_1": "BRL",
                "selling_price_0": "18.40",
                "selling_price_1": "BRL",
                "confirm_lower_price": "",
            },
        )

        self.stock_import.refresh_from_db()

        self.assertEqual(response.status_code, 200)
        self.assertJSONEqual(
            response.headers["HX-Trigger"],
            {
                "productCreated": {},
                "showToast": {
                    "type": "success",
                    "message": "Banana Fila atualizado na importacao manual.",
                },
            },
        )
        self.assertEqual(len(self.stock_import.items_data), 1)
        self.assertEqual(self.stock_import.items_data[0]["qtd"], "3")
        self.assertEqual(self.stock_import.items_data[0]["valor"], "9.10")
        self.assertEqual(self.stock_import.items_data[0]["selling_price"], "18.40")


class StockReportViewTests(TestCase):
    def setUp(self) -> None:
        self.user, self.workshop = create_director_user_with_workshop(suffix=30)
        self.group_filters = CatalogGroup.objects.create(workshop=self.workshop, name="Filtros")
        self.group_engine = CatalogGroup.objects.create(workshop=self.workshop, name="Motor")
        self.supplier_a = Supplier.objects.create(workshop=self.workshop, cnpj="12.345.678/0001-90", name="Fornecedor A")
        self.supplier_b = Supplier.objects.create(workshop=self.workshop, cnpj="98.765.432/0001-10", name="Fornecedor B")

        self.product_filter = Product.objects.create(
            workshop=self.workshop,
            code="FLT-001",
            name="Filtro de Oleo",
            unit=Product.Unit.UND,
            group=self.group_filters,
            cost_price="10.00",
            selling_price="15.00",
            origin_cst=Product.OriginCST.NACIONAL,
            purpose=Product.Purpose.RESALE,
        )
        self.product_oil = Product.objects.create(
            workshop=self.workshop,
            code="OLE-002",
            name="Oleo de Motor",
            unit=Product.Unit.LT,
            group=self.group_filters,
            cost_price="20.00",
            selling_price="30.00",
            origin_cst=Product.OriginCST.NACIONAL,
            purpose=Product.Purpose.RESALE,
        )
        self.product_head = Product.objects.create(
            workshop=self.workshop,
            code="CAB-003",
            name="Cabecote",
            unit=Product.Unit.PC,
            group=self.group_engine,
            cost_price="7.50",
            selling_price="12.00",
            origin_cst=Product.OriginCST.NACIONAL,
            purpose=Product.Purpose.RESALE,
        )

        self.stock_filter = StockProduct.objects.get(workshop=self.workshop, product=self.product_filter)
        self.stock_filter.current_quantity = 5
        self.stock_filter.supplier = self.supplier_a
        self.stock_filter.last_nf = "NF-001"
        self.stock_filter.save(update_fields=["current_quantity", "supplier", "last_nf"])

        self.stock_oil = StockProduct.objects.get(workshop=self.workshop, product=self.product_oil)
        self.stock_oil.current_quantity = 3
        self.stock_oil.supplier = self.supplier_b
        self.stock_oil.last_nf = "NF-002"
        self.stock_oil.save(update_fields=["current_quantity", "supplier", "last_nf"])

        self.stock_head = StockProduct.objects.get(workshop=self.workshop, product=self.product_head)
        self.stock_head.current_quantity = 0
        self.stock_head.save(update_fields=["current_quantity"])

        self.client.force_login(self.user)
        session = self.client.session
        session["active_workshop_id"] = self.workshop.pk
        session.save()

    def test_report_page_renders_default_columns_and_totals(self) -> None:
        response = self.client.get(reverse("stock:report"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Relatorio de Estoque")
        self.assertContains(response, "Exportar Excel")
        self.assertContains(response, "Gerar PDF")
        self.assertContains(response, "FLT-001")
        self.assertContains(response, "Oleo de Motor")
        self.assertContains(response, "R$ 110,00")
        self.assertEqual(response.context["stock_report_totals"]["item_count"], 3)
        self.assertEqual(response.context["stock_report_totals"]["total_quantity"], 8)

    def test_report_page_renders_summary_cards_above_controls_and_export_buttons_on_controls_row(self) -> None:
        response = self.client.get(reverse("stock:report"))

        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        controls_form_index = html.find('id="stock-report-table-controls-form"')
        controls_form_end_index = html.find("</form>", controls_form_index)

        self.assertLess(html.find("Quantidade total"), controls_form_index)
        self.assertLess(html.find('aria-controls="stock-report-table-filter-modal"', controls_form_index, controls_form_end_index), html.find("Exportar Excel", controls_form_index, controls_form_end_index))
        self.assertLess(html.find("Exportar Excel", controls_form_index, controls_form_end_index), html.find("Gerar PDF", controls_form_index, controls_form_end_index))
        self.assertLess(html.find("Gerar PDF", controls_form_index, controls_form_end_index), controls_form_end_index)

    def test_report_filters_by_piece_code_group_supplier_and_quantity_range(self) -> None:
        response = self.client.get(
            reverse("stock:report"),
            data={
                "piece": "Filtro",
                "code": "FLT",
                "group": str(self.group_filters.pk),
                "supplier": str(self.supplier_a.pk),
                "quantity_min": "4",
                "quantity_max": "6",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "FLT-001")
        self.assertNotContains(response, "OLE-002")
        self.assertNotContains(response, "CAB-003")
        self.assertEqual(response.context["stock_report_totals"]["item_count"], 1)
        self.assertEqual(response.context["stock_report_totals"]["total_quantity"], 5)
        self.assertEqual(response.context["stock_report_totals"]["stock_total_cost_display"], "R$ 50,00")

    def test_report_respects_selected_columns_and_preserves_them_in_export_links(self) -> None:
        response = self.client.get(
            reverse("stock:report"),
            data={"columns": ["code", "quantity", "item_total_cost"]},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual([field.label for field in response.context["fields"]], ["Codigo", "Quantidade", "Custo total do item"])
        self.assertContains(response, "columns=code")
        self.assertContains(response, "columns=quantity")
        self.assertContains(response, "columns=item_total_cost")

    def test_report_last_nf_column_shows_friendly_empty_state_text(self) -> None:
        response = self.client.get(
            reverse("stock:report"),
            data={"columns": ["code", "last_nf"]},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "NF-001")
        self.assertContains(response, "NF-002")
        self.assertContains(response, "Nenhuma NF relacionada")

    def test_pdf_preview_view_renders_html_for_iframe(self) -> None:
        response = self.client.get(
            reverse("stock:report_pdf_preview"),
            data={"columns": ["code", "quantity", "item_total_cost"]},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "<!DOCTYPE html>", html=False)
        self.assertContains(response, "Relatorio de Estoque")
        self.assertContains(response, "Custo total do item")
        self.assertIsNone(response.headers.get("X-Frame-Options"))

    @patch("apps.stock.views.build_workshop_logo_data_uri", return_value="data:image/png;base64,bW9uZ28tbG9nbw==")
    def test_pdf_preview_view_renders_workshop_logo(self, build_workshop_logo_data_uri_mock) -> None:
        response = self.client.get(
            reverse("stock:report_pdf_preview"),
            data={"columns": ["code", "quantity", "item_total_cost"]},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'src="data:image/png;base64,bW9uZ28tbG9nbw=="', html=False)
        build_workshop_logo_data_uri_mock.assert_called_once_with(workshop=self.workshop)

    def test_pdf_preview_last_nf_column_shows_friendly_empty_state_text(self) -> None:
        response = self.client.get(
            reverse("stock:report_pdf_preview"),
            data={"columns": ["code", "last_nf"]},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "NF-001")
        self.assertContains(response, "NF-002")
        self.assertContains(response, "Nenhuma NF relacionada")

    @patch("apps.stock.views.render_stock_report_pdf_document")
    def test_pdf_view_returns_downloadable_document_with_filtered_context(self, render_document_mock) -> None:
        render_document_mock.return_value = DocumentPayload(content=b"%PDF-stock", filename="relatorio_estoque.pdf")

        response = self.client.get(
            reverse("stock:report_pdf"),
            data={
                "code": "FLT-001",
                "columns": ["code", "quantity"],
                "download": "1",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, b"%PDF-stock")
        self.assertIn('attachment; filename="relatorio_estoque.pdf"', response["Content-Disposition"])
        context = render_document_mock.call_args.kwargs["context"]
        self.assertEqual(context["stock_report_totals"]["item_count"], 1)
        self.assertEqual([column.label for column in context["selected_columns"]], ["Codigo", "Quantidade"])
        self.assertEqual(context["stock_report_items"][0].product.code, "FLT-001")

    def test_excel_view_returns_workbook_with_selected_columns_and_totals(self) -> None:
        response = self.client.get(
            reverse("stock:report_excel"),
            data={
                "group": str(self.group_filters.pk),
                "quantity_min": "1",
                "columns": ["code", "quantity", "unit_cost", "item_total_cost"],
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn('attachment; filename="relatorio_estoque_', response["Content-Disposition"])

        workbook = load_workbook(filename=BytesIO(response.content))
        sheet = workbook["Relatorio"]

        headers = [sheet.cell(row=10, column=column_index).value for column_index in range(1, 5)]
        self.assertEqual(headers, ["Codigo", "Quantidade", "Custo unitario", "Custo total do item"])
        self.assertEqual(sheet["B6"].value, "2")
        self.assertEqual(sheet["B7"].value, "8")
        self.assertEqual(sheet["B8"].value, "R$ 110,00")
        self.assertEqual(sheet["A11"].value, "FLT-001")
        self.assertEqual(sheet["B11"].value, 5)
        self.assertEqual(sheet["C11"].value, 10)
        self.assertEqual(sheet["D12"].value, 60)

    def test_excel_view_last_nf_column_shows_friendly_empty_state_text(self) -> None:
        response = self.client.get(
            reverse("stock:report_excel"),
            data={"columns": ["code", "last_nf"]},
        )

        self.assertEqual(response.status_code, 200)

        workbook = load_workbook(filename=BytesIO(response.content))
        sheet = workbook["Relatorio"]
        values_by_code = {str(code): value for code, value in sheet.iter_rows(min_row=11, max_col=2, values_only=True) if code is not None}

        self.assertEqual(values_by_code["CAB-003"], "Nenhuma NF relacionada")
        self.assertEqual(values_by_code["FLT-001"], "NF-001")
        self.assertEqual(values_by_code["OLE-002"], "NF-002")


class StockImportSupplierSyncTests(TestCase):
    def setUp(self) -> None:
        self.user, self.workshop = create_director_user_with_workshop(suffix=40)
        self.group = CatalogGroup.objects.create(workshop=self.workshop, name="Filtros")
        self.supplier_old = Supplier.objects.create(workshop=self.workshop, cnpj="12.345.678/0001-40", name="Fornecedor Antigo")
        self.supplier_new = Supplier.objects.create(workshop=self.workshop, cnpj="98.765.432/0001-40", name="Fornecedor Novo")
        self.product = Product.objects.create(
            workshop=self.workshop,
            code="FLT-040",
            name="Filtro Premium",
            unit=Product.Unit.UND,
            group=self.group,
            cost_price="10.00",
            selling_price="15.00",
            origin_cst=Product.OriginCST.NACIONAL,
            purpose=Product.Purpose.RESALE,
        )
        self.stock_product = StockProduct.objects.get(workshop=self.workshop, product=self.product)
        self.stock_product.current_quantity = 2
        self.stock_product.supplier = self.supplier_old
        self.stock_product.last_nf = "NF-OLD"
        self.stock_product.save(update_fields=["current_quantity", "supplier", "last_nf"])

        self.client.force_login(self.user)
        session = self.client.session
        session["active_workshop_id"] = self.workshop.pk
        session.save()

    def test_completed_import_updates_existing_stock_product_supplier_to_latest(self) -> None:
        stock_import = StockImport.objects.create(
            workshop=self.workshop,
            user=self.user,
            nf_key="4" * 44,
            nf_number="NF-NEW",
            supplier_name=self.supplier_new.name,
            supplier_cnpj=self.supplier_new.cnpj,
            items_data=[{"linked_product_id": str(self.product.pk), "qtd": "3"}],
            payments_data=[],
        )
        form = ImportStepSummaryForm(data={}, instance=stock_import, workshop=self.workshop, request=SimpleNamespace(user=self.user))

        self.assertTrue(form.is_valid(), form.errors)
        form.save()

        stock_import.refresh_from_db()
        self.stock_product.refresh_from_db()
        movement = StockMovement.objects.get(stock_product=self.stock_product, type=StockMovement.MovementType.ENTRY, status=StockMovement.MovementStatus.APPROVED)

        self.assertEqual(stock_import.status, StockImport.ImportStatus.COMPLETED)
        self.assertEqual(self.stock_product.current_quantity, 5)
        self.assertEqual(self.stock_product.last_nf, "NF-NEW")
        self.assertEqual(self.stock_product.supplier, self.supplier_new)
        self.assertEqual(movement.supplier, self.supplier_new)


class StockImportPaymentFlowTests(TestCase):
    def setUp(self) -> None:
        self.user, self.workshop = create_director_user_with_workshop(suffix=41)
        self.client.force_login(self.user)
        session = self.client.session
        session["active_workshop_id"] = self.workshop.pk
        session.save()

    def test_add_payment_session_keeps_entered_amount_as_total_paid(self) -> None:
        payment_method = PaymentMethod.objects.create(workshop=self.workshop, description="Cartao", installments_count=4)
        stock_import = StockImport.objects.create(
            workshop=self.workshop,
            user=self.user,
            nf_key="5" * 44,
            items_data=[{"valor": "100.00", "qtd": "1"}],
            payments_data=[],
        )

        response = self.client.post(
            f"{reverse('stock:add_payment_session')}?pk={stock_import.pk}",
            data={
                "payment_method": str(payment_method.pk),
                "payment_date": "2026-03-24",
                "first_amount_0": "40.00",
            },
        )

        stock_import.refresh_from_db()
        self.assertEqual(response.status_code, 204)
        self.assertEqual(len(stock_import.payments_data), 1)
        self.assertEqual(stock_import.payments_data[0]["installments"], "4")
        self.assertEqual(stock_import.payments_data[0]["first_amount"], "40.00")
        self.assertEqual(stock_import.payments_data[0]["total_paid"], "40.00")

    def test_add_additional_value_modal_renders_fields(self) -> None:
        stock_import = StockImport.objects.create(
            workshop=self.workshop,
            user=self.user,
            nf_key="8" * 44,
            items_data=[{"valor": "100.00", "qtd": "1"}],
            payments_data=[],
        )

        response = self.client.get(f"{reverse('stock:add_additional_value_modal')}?pk={stock_import.pk}")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Adicionar Valor")
        self.assertContains(response, "Motivo")
        self.assertContains(response, "Confirmar")

    def test_add_additional_value_session_appends_positive_entry_and_reason(self) -> None:
        stock_import = StockImport.objects.create(
            workshop=self.workshop,
            user=self.user,
            nf_key="9" * 44,
            items_data=[{"valor": "100.00", "qtd": "1"}],
            payments_data=[],
        )

        response = self.client.post(
            f"{reverse('stock:add_additional_value_session')}?pk={stock_import.pk}",
            data={
                "amount_0": "15.00",
                "amount_1": "BRL",
                "reason": "Frete da transportadora",
            },
        )

        stock_import.refresh_from_db()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(stock_import.payments_data), 1)
        self.assertEqual(stock_import.payments_data[0]["entry_type"], "additional_charge")
        self.assertEqual(stock_import.payments_data[0]["amount"], "15.00")
        self.assertEqual(stock_import.payments_data[0]["reason"], "Frete da transportadora")

    def test_add_payment_session_allows_pending_amount_including_additional_values(self) -> None:
        payment_method = PaymentMethod.objects.create(workshop=self.workshop, description="Cartao", installments_count=2)
        stock_import = StockImport.objects.create(
            workshop=self.workshop,
            user=self.user,
            nf_key="1" * 44,
            items_data=[{"valor": "100.00", "qtd": "1"}],
            payments_data=[{"id": 1, "entry_type": "additional_charge", "amount": "20.00", "reason": "Frete"}],
        )

        response = self.client.post(
            f"{reverse('stock:add_payment_session')}?pk={stock_import.pk}",
            data={
                "payment_method": str(payment_method.pk),
                "payment_date": "2026-03-24",
                "first_amount_0": "110.00",
            },
        )

        stock_import.refresh_from_db()
        self.assertEqual(response.status_code, 204)
        self.assertEqual(len(stock_import.payments_data), 2)
        self.assertEqual(stock_import.payments_data[1]["entry_type"], "payment")
        self.assertEqual(stock_import.payments_data[1]["total_paid"], "110.00")

    def test_payment_form_renders_workorder_style_labels_and_table(self) -> None:
        payment_method = PaymentMethod.objects.create(workshop=self.workshop, description="Pix", installments_count=1)
        stock_import = StockImport.objects.create(
            workshop=self.workshop,
            user=self.user,
            nf_key="6" * 44,
            items_data=[{"valor": "100.00", "qtd": "1"}],
            payments_data=[
                {
                    "id": 1,
                    "method": payment_method.pk,
                    "method_display": payment_method.description,
                    "installments": "1",
                    "first_amount": "30.00",
                    "total_paid": "30.00",
                    "payment_date": "2026-03-24",
                    "reason": "Pix",
                },
                {
                    "id": 2,
                    "entry_type": "additional_charge",
                    "amount": "15.00",
                    "reason": "Frete da transportadora",
                },
            ],
        )

        form = ImportStepPaymentForm(instance=stock_import, workshop=self.workshop, import_items=stock_import.items_data, import_payments=stock_import.payments_data)
        html = Template("{% load crispy_forms_tags %}{% crispy form %}").render(Context({"form": form}))

        self.assertIn("Valor Pago", html)
        self.assertIn("Valor Pendente", html)
        self.assertIn("Salvar Plano de Pagamento", html)
        self.assertIn("Adicionar Valor", html)
        self.assertIn("Valor", html)
        self.assertIn("Vencimento", html)
        self.assertIn("Motivo", html)
        self.assertIn("- R$", html)
        self.assertIn("+ R$", html)
        self.assertIn("30,00", html)
        self.assertIn("15,00", html)
        self.assertIn("Frete da transportadora", html)
        self.assertIn("alert_confirm_modal", html)
        self.assertIn('data-confirm="Deseja remover este lançamento financeiro?"', html)

    def test_summary_save_persists_total_without_splitting_installments(self) -> None:
        payment_method = PaymentMethod.objects.create(workshop=self.workshop, description="Cartao", installments_count=4)
        stock_import = StockImport.objects.create(
            workshop=self.workshop,
            user=self.user,
            nf_key="7" * 44,
            nf_number="NF-700",
            items_data=[],
            payments_data=[
                {
                    "id": 1,
                    "method": payment_method.pk,
                    "method_display": payment_method.description,
                    "installments": "4",
                    "first_amount": "40.00",
                    "total_paid": "40.00",
                    "payment_date": "2026-03-24",
                    "reason": "Cartao",
                },
                {
                    "id": 2,
                    "entry_type": "additional_charge",
                    "amount": "15.00",
                    "reason": "Frete",
                },
            ],
            status=StockImport.ImportStatus.DRAFT,
        )

        form = ImportStepSummaryForm(data={}, instance=stock_import, workshop=self.workshop, request=SimpleNamespace(user=self.user))

        self.assertTrue(form.is_valid(), form.errors)
        form.save()

        payment = StockPaymentMethod.objects.get(workshop=self.workshop, nf_number="NF-700")
        self.assertEqual(payment.installments_count, 4)
        self.assertEqual(payment.first_installment_amount, Money("40.00", "BRL"))
        self.assertEqual(payment.remaining_installments_amount, Money("0.00", "BRL"))
        self.assertEqual(payment.total_paid, Money("40.00", "BRL"))
        self.assertEqual(StockPaymentMethod.objects.filter(workshop=self.workshop).count(), 1)


class BackfillStockProductSuppliersCommandTests(TestCase):
    def setUp(self) -> None:
        self.user, self.workshop = create_director_user_with_workshop(suffix=50)
        self.group = CatalogGroup.objects.create(workshop=self.workshop, name="Filtros")
        self.supplier_a = Supplier.objects.create(workshop=self.workshop, cnpj="12.345.678/0001-50", name="Fornecedor A")
        self.supplier_b = Supplier.objects.create(workshop=self.workshop, cnpj="98.765.432/0001-50", name="Fornecedor B")
        self.product = Product.objects.create(
            workshop=self.workshop,
            code="FLT-050",
            name="Filtro Backfill",
            unit=Product.Unit.UND,
            group=self.group,
            cost_price="10.00",
            selling_price="15.00",
            origin_cst=Product.OriginCST.NACIONAL,
            purpose=Product.Purpose.RESALE,
        )
        self.stock_product = StockProduct.objects.get(workshop=self.workshop, product=self.product)

    def test_command_sets_supplier_from_latest_approved_entry_movement(self) -> None:
        StockMovement.objects.create(
            workshop=self.workshop,
            stock_product=self.stock_product,
            type=StockMovement.MovementType.ENTRY,
            supplier=self.supplier_a,
            quantity=1,
            status=StockMovement.MovementStatus.APPROVED,
            transcation_by=self.user,
        )
        StockMovement.objects.create(
            workshop=self.workshop,
            stock_product=self.stock_product,
            type=StockMovement.MovementType.ENTRY,
            supplier=self.supplier_b,
            quantity=1,
            status=StockMovement.MovementStatus.APPROVED,
            transcation_by=self.user,
        )
        StockMovement.objects.create(
            workshop=self.workshop,
            stock_product=self.stock_product,
            type=StockMovement.MovementType.EXIT,
            quantity=1,
            status=StockMovement.MovementStatus.APPROVED,
            transcation_by=self.user,
        )

        call_command("backfill_stock_product_suppliers", workshop_id=self.workshop.pk)

        self.stock_product.refresh_from_db()
        self.assertEqual(self.stock_product.supplier, self.supplier_b)

    def test_command_dry_run_does_not_persist_changes(self) -> None:
        StockMovement.objects.create(
            workshop=self.workshop,
            stock_product=self.stock_product,
            type=StockMovement.MovementType.ENTRY,
            supplier=self.supplier_a,
            quantity=1,
            status=StockMovement.MovementStatus.APPROVED,
            transcation_by=self.user,
        )

        call_command("backfill_stock_product_suppliers", workshop_id=self.workshop.pk, dry_run=True)

        self.stock_product.refresh_from_db()
        self.assertIsNone(self.stock_product.supplier)
